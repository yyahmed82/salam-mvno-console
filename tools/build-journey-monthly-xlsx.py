#!/usr/bin/env python3
"""Monthly journey-errors workbooks (eligibility / BSS activation / Nafath) from the JSONs
that journeyErrExport.js wrote on 152.

  scp -r root@172.31.38.152:/apps/console/server/out/journey-errors ./journey-errors
  python3 tools/build-journey-monthly-xlsx.py ./journey-errors [outdir]

One workbook per month: Salam-Journey-Errors-YYYY-MM.xlsx with tabs
  Summary        month totals per family + the APIGW caveat + activation FK coverage
  Orders         per-lane funnel (created / ID submitted / eligible / denied / activated)
  Eligibility    code × class counts        · Elig rows      failure detail (masked)
  Activation     api × code × class counts  · Act rows       failure detail (masked)
  Nafath         status × service           · Nafath rows    failure detail
Classification follows the console doctrine (salam-issue-classification): Nafath terminal
statuses are business; eligibility/activation split via errclass (605 family = business,
715/5002/1500 = technical). Same house style as the UPG monthlies."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEAD_FILL = PatternFill('solid', fgColor='0B3B2E')
BAND_FILL = PatternFill('solid', fgColor='F4F8F6')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=11)
CLS_COLOR = {'technical': 'B3261E', 'business': '2563EB'}

def sheet(wb, name, header, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for i, r in enumerate(rows):
        ws.append(r)
        if i % 2:
            for c in ws[i + 2]:
                c.fill = BAND_FILL
    for j, w in enumerate(widths or []):
        ws.column_dimensions[chr(65 + j)].width = w
    ws.freeze_panes = 'A2'
    return ws

def build(src, out_dir):
    files = sorted(f for f in os.listdir(src) if f.endswith('.json'))
    if not files:
        sys.exit(f'no JSON files in {src}')
    for f in files:
        d = json.load(open(os.path.join(src, f)))
        m = d['month']
        wb = Workbook(); wb.remove(wb.active)

        et, at = d.get('eligibility_total', {}), d.get('activation_total', {})
        naf_bad = sum(r['n'] for r in d.get('nafath', [])
                      if r['status'] in ('expired', 'rejected', 'failed', 'cancelled', 'denied'))
        fk = (100.0 * at.get('with_fk', 0) / at['n']) if at.get('n') else 0.0
        sheet(wb, 'Summary', ['field', 'value'], [
            ['Month', m], ['Window (UTC)', f"{d['from']} → {d['to']}"],
            ['Eligibility checks / failures', f"{et.get('n', 0):,} / {et.get('bad', 0):,}"],
            ['Activation calls / failures', f"{at.get('n', 0):,} / {at.get('bad', 0):,}"],
            ['Nafath verifications failed', f'{naf_bad:,}'],
            ['Activation→order FK coverage', f'{fk:.1f}% (lane split trustworthy only near 100%)'],
            ['APIGW / Digital-API errors', 'NOT AVAILABLE historically — capture exists since '
             '13 Aug 2026 with 7-day retention. Forward collection only.'],
            ['Classification doctrine', 'technical = Salam IT at fault; Nafath terminal statuses '
             'are ALWAYS business (customer/CITC outcome).'],
            ['Generated', d.get('generated', '')], ['PII', 'masked at extraction'],
        ], [34, 80])

        sheet(wb, 'Orders', ['lane', 'created', 'id_submitted', 'eligible', 'denied', 'activated'],
              [['MNP' if o['mnp'] else 'New SIM', o['created'], o['id_submitted'],
                o['eligible'], o['denied'], o['activated']] for o in d.get('orders', [])],
              [12, 12, 14, 12, 12, 12])

        ws = sheet(wb, 'Eligibility', ['code', 'class', 'count'],
                   [[r['code'], r['cls'], r['n']] for r in d.get('eligibility_codes', [])], [16, 12, 12])
        for row in ws.iter_rows(min_row=2):
            if row[1].value in CLS_COLOR:
                row[1].font = Font(color=CLS_COLOR[row[1].value], bold=True)
        sheet(wb, 'Elig rows', ['created_at', 'api', 'code', 'class', 'response (masked)'],
              [[r['created_at'], r.get('api', ''), r['code'], r['cls'], r.get('response', '')]
               for r in d.get('eligibility_rows', [])], [20, 28, 10, 11, 70])

        ws = sheet(wb, 'Activation', ['api', 'code', 'class', 'count'],
                   [[r['api'], r['code'], r['cls'], r['n']] for r in d.get('activation_codes', [])],
                   [34, 12, 12, 10])
        for row in ws.iter_rows(min_row=2):
            if row[2].value in CLS_COLOR:
                row[2].font = Font(color=CLS_COLOR[row[2].value], bold=True)
        sheet(wb, 'Act rows', ['created_at', 'api', 'code', 'class', 'message (masked)', 'platform', 'has_order_fk'],
              [[r['created_at'], r['api'], r['code'], r['cls'], r.get('message', ''),
                r.get('platform', ''), r.get('has_order', '')] for r in d.get('activation_rows', [])],
              [20, 30, 10, 11, 56, 10, 12])

        sheet(wb, 'Nafath', ['status', 'service', 'count', 'class'],
              [[r['status'], r['service'], r['n'], 'business'] for r in d.get('nafath', [])],
              [14, 20, 10, 11])
        sheet(wb, 'Nafath rows', ['created_at', 'status', 'service'],
              [[r['created_at'], r['status'], r['service']] for r in d.get('nafath_rows', [])],
              [20, 14, 22])

        out = os.path.join(out_dir, f'Salam-Journey-Errors-{m}.xlsx')
        wb.save(out)
        print(f'{m}: elig {len(d.get("eligibility_rows", [])):>6} · '
              f'act {len(d.get("activation_rows", [])):>6} · '
              f'naf {len(d.get("nafath_rows", [])):>5} failure rows → {out}')

if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else './journey-errors'
    out = sys.argv[2] if len(sys.argv) > 2 else '.'
    os.makedirs(out, exist_ok=True)
    build(src, out)
