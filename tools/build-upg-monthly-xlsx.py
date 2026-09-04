#!/usr/bin/env python3
"""Build one workbook per month from the UPG export, matching Payment-Failures-Detail-20Aug2026.xlsx.

    python3 tools/build-upg-monthly-xlsx.py <indir> <outdir>

<indir> is the folder produced by server/src/upgMonthlyExport.js on 152:
    YYYY-MM-failures.csv · YYYY-MM-attempts.csv · YYYY-MM-voucher.json · YYYY-MM-meta.json

Each workbook carries the same six tabs, in the same order, with the same headers, header style
(white bold on #0E9F5A), frozen panes, autofilters and column widths as the 20 Aug file:
    README · Categories + TAP samples · Voucher recharge · Summary · Failed payments · Gateway attempts

TWO RULES THAT LOOK LIKE DETAILS AND ARE NOT

1. CATEGORY GROUPING. Categories group by (CLASS, bank_message) — so "Abandoned" appears once with
   its rails joined ("CARD / STC_PAY"). The ONE exception is an empty bank_message, which groups by
   (CLASS, message, rail): on a blank message the rail is the entire distinguishing fact, and a
   wallet blank (by design) must never be merged with a card blank (a real gap). This reproduces the
   20 Aug sheet exactly, including why STC_PAY blanks appear on two rows: 1,691 BUSINESS + 1
   RECONCILIATION = the 1,692 in the Summary.

2. THE VOUCHER TAB IS NOT BACKFILLABLE. Voucher recharges write no DB row; the only record is the
   console API capture, which has 7-day retention and began on 13 Aug 2026. For every month before
   that the tab states plainly that nothing was retained. It is never estimated and the August
   figures are never carried backwards.
"""
import csv, json, os, sys, glob, collections, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IN  = sys.argv[1] if len(sys.argv) > 1 else 'upg'
OUT = sys.argv[2] if len(sys.argv) > 2 else '.'

HDR_FILL = PatternFill('solid', fgColor='0E9F5A')
HDR_FONT = Font(bold=True, color='FFFFFF')
TITLE    = Font(bold=True, size=13)

# ---------------------------------------------------------------- category reference text
# Keyed on the acquirer message. Anything not listed still renders, with a neutral reading — a new
# message must never be silently dropped from the sheet just because nobody has described it yet.
MEAN = {
 'Abandoned': ('Payment page opened, customer never entered/confirmed payment. TAP auto-expires the '
               'charge ~31 min later.', 'Customer behaviour — recover via reminder journeys'),
 '(no acquirer message)': ('Wallet rail returns no acquirer message — BY DESIGN. Successful Apple Pay / '
               'STC Pay payments carry a blank message too.', 'No action — reporting semantics'),
 'Declined, Transaction Type not Supported': ('Acquirer refused this transaction TYPE on the MADA rail.',
               'ACTION: UPG MS to verify transactionType enabled for MADA on our MID'),
 'Declined, Insufficient Funds': ('Bank refused: not enough balance on the card.', 'Customer/bank'),
 'Declined, Incorrect CSC/CVV': ('Wrong CVV entered.', 'Customer'),
 'Failed': ('Generic gateway failure with no acquirer detail — the only technical category with volume.',
            'ACTION: UPG/TAP to clarify what generic "Failed" means on STC_PAY'),
 'Failed, Expired Card': ('Card expiry date has passed.', 'Customer'),
 'Declined, Card Issuer': ('Issuing bank declined without a specific reason code.', 'Customer/bank'),
 'Declined': ('Generic bank decline (no sub-reason returned).', 'Customer/bank'),
 'Declined, Authentication Failed': ('3-D Secure / OTP authentication failed.', 'Customer/bank'),
 'Declined, Not Authenticated': ('3-D Secure authentication not completed.', 'Customer/bank'),
 'Restricted, Bank': ('Card restricted by the issuing bank.', 'Customer/bank'),
 'Failed, Invalid Card No': ('Card number invalid (typo / wrong PAN).', 'Customer'),
 'Failed, Unspecified Failure': ('Bank returned a failure without a reason code.', 'Customer/bank'),
 'Restricted': ('Card restricted.', 'Customer/bank'),
 'Declined, Card Issuer - Error': ('Issuer returned an error response.', 'Customer/bank'),
 'Timed Out': ('Gateway timed out before an answer.', 'Platform'),
 'Declined, Card Issuer - Referral Response': ('Issuer asked for referral (manual authorisation).',
                                               'Customer/bank'),
 'Captured': ('—', '—'),
}
NO_MSG = '(no acquirer message)'
VOUCHER_NOTE = {'22': 'Voucher already consumed', '21': 'Wrong number entered',
                '0': 'CODE "0" vs success "00" — identical if read as a number',
                '9': 'Policy block', '31': 'Platform/service error'}


def style_header(ws, row=1):
    for c in ws[row]:
        if c.value is not None:
            c.fill, c.font = HDR_FILL, HDR_FONT


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def read_csv(p):
    if not os.path.exists(p): return []
    with open(p, newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


def month_label(ym):
    return datetime.date(int(ym[:4]), int(ym[5:7]), 1).strftime('%B %Y')


# ---------------------------------------------------------------- one workbook
def build(ym, indir, outdir):
    fails = read_csv(os.path.join(indir, f'{ym}-failures.csv'))
    atts  = read_csv(os.path.join(indir, f'{ym}-attempts.csv'))
    if not fails:
        print(f'  {ym}  no failures csv — skipped'); return None
    vpath = os.path.join(indir, f'{ym}-voucher.json')
    vouch = json.load(open(vpath)) if os.path.exists(vpath) else {'retained': False, 'rows': [],
             'note': 'voucher export not present for this month'}
    meta = {}
    mpath = os.path.join(indir, f'{ym}-meta.json')
    if os.path.exists(mpath): meta = json.load(open(mpath))

    n = len(fails)
    label = month_label(ym)
    pc = lambda x: round(x * 100.0 / n, 2) if n else 0

    wb = Workbook(); wb.remove(wb.active)
    # masking state comes from the EXPORT's own meta — never assumed (1 Sep 2026: business asked
    # for unmasked monthlies; a workbook must state which kind it is, in its title AND filename)
    masked = bool(meta.get('pii_masked', True))

    # ---- README
    ws = wb.create_sheet('README')
    rows = [
     ('Payment Failures — per-transaction detail' + ('' if masked else '  ·  INTERNAL — UNMASKED'), ''),
     ('Period', f'{label} (created_at UTC)'),
     ('Scope', f'Every payment the app marked failed: {n:,} rows, each resolved against the UPG gateway'),
     ('Sheets', '"Failed payments" = one row per failed payment (final gateway outcome) · '
                '"Gateway attempts" = every individual gateway attempt on those references (retries)'),
     ('CLASSIFICATION RULE (agreed 20 Aug 2026)', ''),
     ('BUSINESS', 'Customer or bank outcome — not a platform fault. Includes: bank declines (funds, CVV, '
                  'expired, issuer), customer abandonment (page opened, never engaged), and wallet-rail '
                  'transactions with no acquirer message (Apple Pay / STC Pay return none by design — '
                  'their successful payments are blank too).'),
     ('TECHNICAL', 'Platform fault only (gateway/transport failures, timeouts, unspecified failures).'),
     ('RECONCILIATION', 'App marked failed while the gateway shows the payment as PAID.'),
     ('COLUMN GUIDE', ''),
     ('reference_id', 'Our payment reference = the gateway invoice id — the join key between both sides.'),
     ('rail / method', 'The real payment rail from the gateway record (CARD / APPLE_PAY / STC_PAY) and the '
                       'scheme (MADA/VISA/…). NOTE: the app column card_type is hardcoded and must not be used.'),
     ('bank_message', 'The acquirer/bank answer. Empty on wallet rails by design.'),
     ('gateway_attempts', 'How many times this reference was attempted at the gateway (customer retries).'),
     ('ever_paid_at_gateway', 'yes = a later attempt on the same reference succeeded.'),
     ('customer_mobile / target_mobile',
      'MASKED. An internal unmasked copy can be produced on request (audited).' if masked else
      'UNMASKED — FULL MOBILE NUMBERS. Internal distribution only: do not forward outside Salam, '
      'do not attach to external tickets. Produced on business request; the generation is audited.'),
     ('Voucher recharge tab',
      'Voucher recharges write no database row (BSS is called directly). The only record is the console '
      'API capture, which has 7-day retention and started 13 Aug 2026 — so this tab can only be '
      'populated for recent months. Where it is empty, the data was never retained; it is not estimated.'),
     ('Source', 'Salam Digital Console — app replica joined to the UPG gateway on reference id '
                '(read-only). Reproducible on demand for any period.'),
    ]
    for r in rows: ws.append(r)
    ws['A1'].font = Font(bold=True, size=14)
    for r in (5, 9): ws.cell(r, 1).font = Font(bold=True)
    widths(ws, {'A': 34, 'B': 118})
    for row in ws.iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical='top')

    # ---- Categories + TAP samples
    ws = wb.create_sheet('Categories + TAP samples')
    ws.append([f'Payment error categories — {label} — with real TAP charge references'])
    ws.append(['Every category below is a real bank/gateway answer. TAP sample columns carry actual '
               'charge ids so any row can be verified at the gateway.'])
    ws.append([])
    ws.append(['CLASS', 'Category (bank message)', 'Rail', 'Count', '% of failed', 'What it means',
               'Owner / action', 'TAP sample 1 (charge id)', 'sample 1 · when (UTC)',
               'sample 1 · amount SAR', 'sample 1 · our reference', 'TAP sample 2 (charge id)',
               'TAP sample 3 (charge id)'])
    ws['A1'].font = TITLE
    style_header(ws, 4)

    groups = collections.OrderedDict()
    for r in fails:
        msg = (r['bank_message'] or '').strip()
        # see rule 1 in the module docstring: blanks group by rail, everything else merges rails
        key = (r['CLASS'], msg or NO_MSG, r['rail'] if not msg else None)
        g = groups.setdefault(key, {'n': 0, 'rails': set(), 'samples': []})
        g['n'] += 1
        g['rails'].add(r['rail'] or '?')
        if len(g['samples']) < 3 and r['gateway_transaction_id']:
            g['samples'].append(r)
    for (cls, msg, _rail), g in sorted(groups.items(), key=lambda kv: -kv[1]['n']):
        mean, owner = MEAN.get(msg, ('—', '—'))
        s = g['samples']
        ws.append([cls, msg, ' / '.join(sorted(g['rails'])), g['n'], pc(g['n']), mean, owner,
                   s[0]['gateway_transaction_id'] if s else '',
                   s[0]['created_at_utc'] if s else '',
                   s[0]['amount_sar'] if s else '',
                   s[0]['reference_id'] if s else '',
                   s[1]['gateway_transaction_id'] if len(s) > 1 else '',
                   s[2]['gateway_transaction_id'] if len(s) > 2 else ''])
    by_class = collections.Counter(r['CLASS'] for r in fails)
    ws.append([])
    for c in ('BUSINESS', 'TECHNICAL', 'RECONCILIATION'):
        if by_class.get(c):
            ws.append([f'TOTAL {c}', '', '', by_class[c], pc(by_class[c])])
    ws.freeze_panes = 'A5'
    widths(ws, {'A': 13, 'B': 42, 'C': 20, 'D': 9, 'E': 11, 'F': 64, 'G': 44, 'H': 30, 'I': 19,
                'J': 13, 'K': 15, 'L': 30, 'M': 30})

    # ---- Voucher recharge
    ws = wb.create_sheet('Voucher recharge')
    if vouch.get('retained') and vouch.get('rows'):
        vr = vouch['rows']
        tot = sum(int(r['n']) for r in vr)
        succ = sum(int(r['n']) for r in vr if str(r['code']) == '00')
        failed = tot - succ
        ws.append([f'Digital voucher recharge — response codes, {label}'])
        ws.append([f'{tot:,} attempts captured (rolling 7-day capture window — see README) · '
                   'classification applied to the response-code list confirmed by the platform team'])
        ws.append([])
        ws.append(['CLASS', 'Code', 'Response message', 'Count', '% of attempts', '% of failures',
                   'Avg ms', 'Max ms', 'Note'])
        style_header(ws, 4)
        for r in sorted(vr, key=lambda x: -int(x['n'])):
            code, cnt = str(r['code']), int(r['n'])
            is_succ = code == '00'
            cls = 'SUCCESS' if is_succ else ('TECHNICAL' if code in ('0', '31') else 'BUSINESS')
            ws.append([cls, code, r['msg'], cnt, round(cnt * 100 / tot, 2) if tot else 0,
                       '' if is_succ else (round(cnt * 100 / failed, 2) if failed else 0),
                       num(r.get('avg_ms')), num(r.get('max_ms')), VOUCHER_NOTE.get(code, '')])
        biz = sum(int(r['n']) for r in vr
                  if str(r['code']) != '00' and str(r['code']) not in ('0', '31'))
        tech = failed - biz
        ws.append(['TOTALS', '', 'attempts', tot, 100])
        ws.append(['', '', 'failed', failed, round(failed * 100 / tot, 2) if tot else 0, 100])
        ws.append(['', '', '  · business', biz, '', round(biz * 100 / failed, 1) if failed else 0,
                   '', '', 'wrong / used / reserved / blocked voucher'])
        ws.append(['', '', '  · technical', tech, '', round(tech * 100 / failed, 1) if failed else 0,
                   '', '', 'FailedInBRM + service error 31'])
    else:
        ws.append([f'Digital voucher recharge — {label}'])
        ws.append(['NO DATA RETAINED FOR THIS PERIOD — this is a retention fact, not a zero.'])
        ws.append([])
        ws.append(['Why', vouch.get('note', '')])
        ws.append(['What this does NOT mean', 'It does not mean there were no voucher recharges or no '
                   'voucher failures this month. It means no record of them survives to be counted.'])
        ws.append(['To capture it going forward', 'Voucher attempts are only visible through the console '
                   'API capture (api_traffic_events, 7-day retention). Raising that retention, or '
                   'archiving the voucher lane monthly, is the only way to have this tab for future '
                   'periods.'])
        ws['A1'].font = TITLE
        ws['A2'].font = Font(bold=True, color='B3261E')
        for r in (4, 5, 6): ws.cell(r, 1).font = Font(bold=True)
        for row in ws.iter_rows(min_col=2, max_col=2):
            row[0].alignment = Alignment(wrap_text=True, vertical='top')
    if ws['A4'].value == 'CLASS': ws.freeze_panes = 'A5'
    ws['A1'].font = TITLE
    widths(ws, {'A': 24, 'B': 90, 'C': 46, 'D': 9, 'E': 14, 'F': 13, 'G': 9, 'H': 9, 'I': 52})
    if vouch.get('retained'): widths(ws, {'A': 12, 'B': 7})

    # ---- Summary
    ws = wb.create_sheet('Summary')
    ws.append([f'Payment failures {label} — summary'])
    ws.append(['total failed payments', n])
    ws.append(['gateway attempt records', len(atts)])
    ws['A1'].font = TITLE

    def block(title, counter, denom=None):
        ws.append([])
        ws.append([title, 'count', '% of failed'])
        style_header(ws, ws.max_row)
        for kk, vv in counter.most_common():
            ws.append([kk if kk else '(none)', vv, round(vv * 100.0 / (denom or n), 2)])

    block('BY CLASS', collections.Counter(r['CLASS'] for r in fails))
    block('BY SUB-CLASS', collections.Counter(r['sub_class'] for r in fails))
    block('BY RAIL', collections.Counter(r['rail'] for r in fails))
    block('BY PAYMENT TYPE', collections.Counter(r['payment_type'] for r in fails))
    block('BY PLATFORM', collections.Counter(r['platform'] for r in fails))
    block('TOP BANK MESSAGES · RAIL', collections.Counter(
        f"{(r['bank_message'] or '(no message)')} · {r['rail'] or '?'}" for r in fails))
    widths(ws, {'A': 56, 'B': 12, 'C': 12})

    # ---- Failed payments / Gateway attempts (raw)
    for name, rows_, wspec in [
        ('Failed payments', fails, {'A': 38, 'B': 19, 'C': 15, 'D': 22, 'E': 38, 'F': 9, 'H': 12,
                                    'I': 9, 'J': 11, 'K': 10, 'L': 22, 'M': 14, 'O': 20, 'P': 13,
                                    'Q': 38, 'R': 30}),
        ('Gateway attempts', atts, {'A': 15, 'B': 11, 'C': 20, 'D': 19, 'E': 11, 'F': 12, 'G': 9,
                                    'H': 15, 'I': 38, 'J': 30})]:
        ws = wb.create_sheet(name)
        if not masked and name in ('Failed payments', 'Gateway attempts'):
            ws.append(['⚠ UNMASKED INTERNAL COPY — contains full customer mobile numbers. '
                       'Internal distribution only.'])
            ws['A1'].font = Font(bold=True, color='FFFFFFFF')
            ws['A1'].fill = PatternFill('solid', fgColor='FFDC2626')
        if not rows_:
            ws.append(['(no rows)']); continue
        head = list(rows_[0].keys())
        ws.append(head)
        style_header(ws)
        numeric = {'amount_sar', 'gateway_attempts', 'attempt_no', 'amount_halalas'}
        for r in rows_:
            ws.append([num(r[h]) if h in numeric else r[h] for h in head])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(head))}{len(rows_) + 1}'
        widths(ws, wspec)

    # ---- self-check against the extractor's own count before writing anything
    if meta and meta.get('failures') not in (None, n):
        raise SystemExit(f'{ym}: meta says {meta["failures"]} failures, csv has {n} — refusing to '
                         f'write a workbook whose contents disagree with the extract.')

    out = os.path.join(outdir, f'Salam-Payment-Failures-{ym}' + ('' if masked else '-UNMASKED-INTERNAL') + '.xlsx')
    wb.save(out)
    print(f'  {ym}  {n:>7,} failures · {len(atts):>7,} attempts · voucher '
          f'{"yes" if vouch.get("retained") else "not retained"}  → {os.path.basename(out)}')
    return out


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    months = sorted({os.path.basename(p)[:7] for p in glob.glob(os.path.join(IN, '*-failures.csv'))})
    if not months:
        sys.exit(f'no *-failures.csv in {IN}')
    print(f'building {len(months)} monthly workbooks from {IN}')
    made = [build(m, IN, OUT) for m in months]
    print(f'\n✓ {len([x for x in made if x])} workbooks in {OUT}')
