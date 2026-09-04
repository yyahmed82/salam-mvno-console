#!/usr/bin/env python3
"""HyperPay failures workbook for the urgent STC Pay mail.
Usage: python3 tools/build-hyperpay-failures-xlsx.py /path/hyperpay-failures.json ../HyperPay-Failures-Since-Cutover.xlsx
Sheet 1: summary (by brand, by result code). Sheet 2: every failed transaction row."""
import json, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

d = json.load(open(sys.argv[1])); rows = d["rows"]
wb = Workbook(); ws = wb.active; ws.title = "Summary"
G = PatternFill("solid", fgColor="1F7A4D"); W = Font(bold=True, color="FFFFFF")
ws.append([f"HyperPay customer-payment FAILURES since cutover {d['cutover']} (KSA) · generated {d['generated_at'][:16]}Z"])
ws["A1"].font = Font(bold=True, size=12)
ws.append([]); ws.append(["By brand", "failures"])
for c in ws[3]: c.font = W; c.fill = G
for b, n in sorted(d["by_brand"].items(), key=lambda x: -x[1]): ws.append([b, n])
ws.append([]); ws.append(["By HyperPay result code", "count", "description (sample)"])
for c in ws[ws.max_row]: c.font = W; c.fill = G
cc = Counter((r.get("result_code") or "—", (r.get("result_description") or r.get("fail_reason") or "")[:80]) for r in rows)
for (code, desc), n in cc.most_common(15): ws.append([code, n, desc])
ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 70

ws2 = wb.create_sheet("All failures")
hdr = ["When (KSA)", "Brand", "Platform", "Amount SAR", "HyperPay txn id", "Checkout id",
       "Result code", "Result description", "App fail reason", "Payment type", "Customer mobile", "App payment id"]
ws2.append(hdr)
for c in ws2[1]: c.font = W; c.fill = G
for r in rows:
    ws2.append([str(r.get("at_ksa", ""))[:19], r.get("brand"), r.get("platform"), r.get("amount"),
                r.get("gateway_txn_id"), r.get("checkout_id"), r.get("result_code"),
                r.get("result_description"), r.get("fail_reason"), r.get("payment_on_type"),
                r.get("customer_mobile"), r.get("payment_id")])
for col, w in zip("ABCDEFGHIJKL", [19, 10, 9, 10, 42, 16, 13, 40, 34, 15, 14, 38]):
    ws2.column_dimensions[col].width = w
ws2.auto_filter.ref = f"A1:L{ws2.max_row}"
wb.save(sys.argv[2]); print(f"wrote {sys.argv[2]} · {len(rows)} failures · brands {d['by_brand']}")
