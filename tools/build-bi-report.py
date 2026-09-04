#!/usr/bin/env python3
"""BI PAYMENT REPORT — STEP 3 of 3: build the workbook from the extracted JSON.

Input : bi2/bi/YYYY-MM.json          (biExtract.js)   — funnel, journey, reasons, vendor/rail/platform
        bi2/bi/identity-YYYY-MM.json (biIdentity.js)  — account membership for journeys 3 + 4
Output: Salam-Payment-Errors-Sep2025-Aug2026.xlsx

Structure mirrors Payment-Errors-Classified-Catalog-20Aug2026 so BI can lay them side by side,
plus the four journeys they asked for.

Every caveat that could mislead a reader is on the README sheet, in plain words — not in a
footnote. The three that matter:
  · the window is 12 months, not 20 — payments has a 6-month hole (Mar–Aug 2025)
  · "never_attempted" is 38% of all rows and is NOT an error
  · the recharge/invoice split is ACCOUNT MEMBERSHIP, not session state

Usage:  python3 tools/build-bi-report.py [indir] [outfile]
"""
import json, glob, os, sys, collections

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference

IN  = sys.argv[1] if len(sys.argv) > 1 else 'bi2/bi'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'Salam-Payment-Errors-Sep2025-Aug2026.xlsx'

INK   = '11241D'; GREEN = '00A651'; DARK = '0B3B2E'
LIGHT = 'F2F6F3'; AMBER = 'B26B00'; RED  = 'B3261E'; GREY = '5F6F69'

HDR   = PatternFill('solid', fgColor=DARK)
SUBHD = PatternFill('solid', fgColor=LIGHT)
THIN  = Side(style='thin', color='DDE7E1')
BOX   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

JOURNEY_LABEL = {
    '1_new_sim': '① New SIM',
    '2_mnp': '② MNP port-in',
    '3_recharge': '③ Recharge',
    '4_invoice': '④ Invoice / bill payment',
    '5_checkout': '⑤ Checkout (renewal, plan change, …)',
    '1_2_order_type_unknown': 'Order — number type not set',
    'other': 'Other',
}
OUTCOMES = ['success', 'failed', 'never_attempted', 'stuck_callback', 'refunded']

# ---------------------------------------------------------------- load
months, data, ident = [], {}, {}
for f in sorted(glob.glob(os.path.join(IN, '20*.json'))):
    d = json.load(open(f)); months.append(d['month']); data[d['month']] = d
for f in sorted(glob.glob(os.path.join(IN, 'identity-*.json'))):
    d = json.load(open(f)); ident[d['month']] = d
if not months:
    sys.exit(f'no month files found in {IN}')
print(f'{len(months)} months: {months[0]} → {months[-1]}   ({len(ident)} identity files)')

num = lambda v: float(v) if v not in (None, '') else 0.0
wb = Workbook(); wb.remove(wb.active)


def sheet(title, widths=None):
    ws = wb.create_sheet(title)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def head(ws, row, cols, fill=HDR, color='FFFFFF'):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(bold=True, color=color, size=10)
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, row, values, bold=False, money=(), pct=()):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.border = BOX
        if bold: cell.font = Font(bold=True)
        if isinstance(v, (int, float)):
            cell.number_format = '#,##0.00' if i in money else ('0.0%' if i in pct else '#,##0')


# ---------------------------------------------------------------- 1 · README
ws = sheet('README', [3, 30, 100])
ws['B2'] = 'Salam · Payment errors — classified catalog'
ws['B2'].font = Font(bold=True, size=16, color=DARK)
ws['B3'] = f'{months[0]} → {months[-1]} · 12 months · built {__import__("datetime").date.today()}'
ws['B3'].font = Font(size=11, color=GREY)

rows = [
    ('', ''),
    ('READ THIS FIRST', ''),
    ('Period is 12 months, not 20',
     'BI asked for Jan 2025 onward. The source cannot serve it. The payments table starts '
     '2025-02-14 and then has a SIX-MONTH HOLE — March to August 2025 contain zero rows — before '
     'resuming complete from September 2025. onboarding_orders (from 2022-01) and checkouts '
     '(from 2022-06) have no such gap, so this is specific to payments and is almost certainly '
     'the archival job (the table carries archived / archived_at columns). '
     'Sep 2025 → Aug 2026 is the longest unbroken window that exists.'),
    ('August 2026 is partial',
     'The month was still in progress when the data was taken. Do not read Aug as a full month '
     'or compare its totals to the others.'),
    ('"never_attempted" is NOT an error',
     'It is 38% of all payment rows and it is the single biggest number in this report. It means '
     'the payment record was created but the customer never reached or never finished the payment '
     'page, so no gateway answer exists. It is a FUNNEL GAP, not a failure. Counting it as an '
     'error overstates failure roughly fourfold.'),
    ('"stuck_callback" IS the technical bucket',
     'Pending, but the gateway did answer. Money may have been taken while our record was never '
     'finalised. Small (0.1%) and the one to watch.'),
    ('refunded is reported separately', 'Refunds are neither success nor failure and have their own column.'),
    ('The two Aug 2026 figures differ by ~190',
     'The account-membership tables were extracted a few minutes after the main tables, and August '
     'was still accumulating between the two passes. Sep 2025 – Jul 2026 reconcile EXACTLY, row for '
     'row; only August drifts, by 193 payments out of 213,000 (0.09%). Nothing is missing — it is '
     'two snapshots of a live month taken minutes apart. Any month-on-month analysis should end at '
     'July 2026.'),
    ('', ''),
    ('THE LOGGED / NOT-LOGGED SPLIT', ''),
    ('It is three different things',
     'The app records identity three different ways, so this report reports three different things '
     'rather than pretending to one. Read the label on each sheet.'),
    ('① New SIM and ② MNP: no split',
     'onboarding_orders.orderable_type is 99.99% AnonymousUser. That is CORRECT behaviour — a '
     'customer has no account before buying their first SIM. A logged/guest split here would be '
     'invented, so none is shown.'),
    ('⑤ Checkout: a true split',
     'checkouts.checkoutable_type genuinely carries User / Guest / AnonymousUser / DeletedUser. '
     'This is the only journey where "logged in vs not" is directly recorded.'),
    ('③ Recharge and ④ Invoice: ACCOUNT MEMBERSHIP, NOT SESSION STATE',
     'These payments carry no link to any identity at all — payment_on_id is NULL on 100% of them. '
     'What is reported instead is whether the paying mobile EXISTS in users or in guests. '
     'It tells you whether that number is a registered customer. It does NOT tell you whether the '
     'person was logged in when they paid — the app does not record that. Please do not relabel '
     'these columns "logged in / not logged in".'),
    ('A proxy that was tested and rejected',
     'Comparing customer_mobile_number to target_mobile_number was tried first. Over 12 months it '
     'returned 3,509,083 "same" against 24 "different" — the app writes both columns with the same '
     'value, so the comparison is a constant, not a signal. It was dropped rather than shipped.'),
    ('', ''),
    ('METHOD', ''),
    ('Source', 'Production REPLICA (salam_replica), read-only. No query in this pipeline touched the live database.'),
    ('Load', 'One aggregated pass per month, index-driven, with a pause between months. No row-level extract.'),
    ('Journey mapping',
     '① OnboardingOrder + number_order_type=0 · ② OnboardingOrder + number_order_type=1 · '
     '③ recharge + postpaid_service_recharge · ④ bill + advanced_postpaid_payment · ⑤ Checkout'),
    ('Failure reason',
     'fail_reason first; if empty, "code · message" from the stored gateway response. UPG/salam '
     'leave fail_reason empty and put the truth in the gateway snapshot. Same definition as the '
     'console Troubleshoot board and the 20 Aug catalog.'),
    ('"(no message)"',
     'The gateway returned no reason at all. Largest single category on Apple Pay — raised with '
     'the provider separately.'),
]
r = 5
for k, v in rows:
    if k and not v:
        ws.cell(row=r, column=2, value=k).font = Font(bold=True, size=12, color=GREEN)
    else:
        c = ws.cell(row=r, column=2, value=k); c.font = Font(bold=True, size=10); c.alignment = Alignment(vertical='top')
        c2 = ws.cell(row=r, column=3, value=v); c2.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = max(15, 13 * (len(v) // 95 + 1))
    r += 1

# ---------------------------------------------------------------- 2 · Summary
ws = sheet('Summary', [12] + [14] * 8)
ws['A1'] = 'Monthly funnel — all payments'
ws['A1'].font = Font(bold=True, size=13, color=DARK)
head(ws, 3, ['Month', 'Total', 'Success', 'Failed', 'Never attempted', 'Stuck callback',
             'Success rate', 'Failure rate', 'Success SAR'])
r = 4
tot = collections.Counter()
for m in months:
    f = data[m]['funnel']
    t = f['total']
    body(ws, r, [m, t, f['success'], f['failed'], f['never_attempted'], f['stuck_callback'],
                 f['success'] / t if t else 0, f['failed'] / t if t else 0, num(f['success_sar'])],
         money=(9,), pct=(7, 8))
    for k in ('total', 'success', 'failed', 'never_attempted', 'stuck_callback'): tot[k] += f[k]
    tot['sar'] += num(f['success_sar'])
    r += 1
body(ws, r, ['TOTAL', tot['total'], tot['success'], tot['failed'], tot['never_attempted'],
             tot['stuck_callback'], tot['success'] / tot['total'], tot['failed'] / tot['total'], tot['sar']],
     bold=True, money=(9,), pct=(7, 8))

ch = LineChart(); ch.title = 'Payments by outcome, per month'; ch.height = 9; ch.width = 22
ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=3, max_row=3 + len(months)), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + len(months)))
ws.add_chart(ch, 'K3')

# ---------------------------------------------------------------- 3 · Journeys × month
ws = sheet('Journeys by month', [12, 34] + [14] * 7)
ws['A1'] = 'The four journeys BI asked for (⑤ Checkout included for completeness)'
ws['A1'].font = Font(bold=True, size=13, color=DARK)
head(ws, 3, ['Month', 'Journey', 'Total', 'Success', 'Failed', 'Never attempted',
             'Stuck callback', 'Refunded', 'Success rate'])
r = 4
jt = collections.defaultdict(collections.Counter)
for m in months:
    agg = collections.defaultdict(collections.Counter)
    for row in data[m]['journey']:
        agg[row['journey']][row['outcome']] += row['n']
        agg[row['journey']]['total'] += row['n']
    for j in sorted(agg):
        a = agg[j]
        body(ws, r, [m, JOURNEY_LABEL.get(j, j), a['total'], a['success'], a['failed'],
                     a['never_attempted'], a['stuck_callback'], a['refunded'],
                     a['success'] / a['total'] if a['total'] else 0], pct=(9,))
        for k in a: jt[j][k] += a[k]
        r += 1

# ---------------------------------------------------------------- 4 · per-journey sheets
SHEET_NAME = {                      # Excel forbids / \ ? * [ ] : in sheet titles, and caps at 31 chars
    '1_new_sim': '1 New SIM',
    '2_mnp': '2 MNP port-in',
    '3_recharge': '3 Recharge',
    '4_invoice': '4 Invoice payment',
    '5_checkout': '5 Checkout',
}

def journey_sheet(jkey, identity_rows=None, identity_note=None):
    name = SHEET_NAME[jkey]
    ws = sheet(name, [12] + [16] * 8)
    ws['A1'] = JOURNEY_LABEL[jkey]
    ws['A1'].font = Font(bold=True, size=13, color=DARK)
    if identity_note:
        ws['A2'] = identity_note
        ws['A2'].font = Font(italic=True, size=10, color=AMBER)
    head(ws, 4, ['Month', 'Total', 'Success', 'Failed', 'Never attempted', 'Stuck', 'Refunded', 'Success rate'])
    r = 5
    for m in months:
        a = collections.Counter()
        for row in data[m]['journey']:
            if row['journey'] == jkey:
                a[row['outcome']] += row['n']; a['total'] += row['n']
        body(ws, r, [m, a['total'], a['success'], a['failed'], a['never_attempted'],
                     a['stuck_callback'], a['refunded'], a['success'] / a['total'] if a['total'] else 0], pct=(8,))
        r += 1
    t = jt[jkey]
    body(ws, r, ['TOTAL', t['total'], t['success'], t['failed'], t['never_attempted'],
                 t['stuck_callback'], t['refunded'], t['success'] / t['total'] if t['total'] else 0],
         bold=True, pct=(8,))

    if identity_rows:
        r += 3
        ws.cell(row=r, column=1, value=identity_rows['title']).font = Font(bold=True, size=12, color=DARK)
        r += 1
        ws.cell(row=r, column=1, value=identity_rows['caveat']).font = Font(italic=True, size=9, color=AMBER)
        r += 1
        ws.cell(row=r, column=1,
                value='Aug 2026 totals here are ~190 higher than the table above: extracted minutes later, '
                      'and August was still accumulating. All other months reconcile exactly.'
                ).font = Font(italic=True, size=9, color=GREY)
        r += 2
        head(ws, r, ['Month'] + identity_rows['cols'] + ['Total'])
        r += 1
        for m in months:
            vals = identity_rows['by_month'].get(m, {})
            tt = sum(vals.get(c, 0) for c in identity_rows['cols'])
            body(ws, r, [m] + [vals.get(c, 0) for c in identity_rows['cols']] + [tt])
            r += 1
        allv = identity_rows['total']
        body(ws, r, ['TOTAL'] + [allv.get(c, 0) for c in identity_rows['cols']] +
             [sum(allv.get(c, 0) for c in identity_rows['cols'])], bold=True)
    return ws


ACC_COLS = ['has_user_account', 'guest_record_only', 'no_record', 'no_mobile_recorded']

def identity_block(jkey):
    by_month, total = {}, collections.Counter()
    for m in months:
        c = collections.Counter()
        for row in ident.get(m, {}).get('rows', []):
            if row['journey'] == jkey:
                c[row['account_status']] += row['n']; total[row['account_status']] += row['n']
        by_month[m] = c
    return {
        'title': 'Account membership of the paying mobile',
        'caveat': 'NOT session state. This says whether the number is registered in users / guests — '
                  'not whether the person was logged in when they paid. The app does not record that.',
        'cols': ACC_COLS, 'by_month': by_month, 'total': total,
    }

journey_sheet('1_new_sim',
              identity_note='No logged/guest split: 99.99% of onboarding orders are AnonymousUser — '
                            'customers have no account before buying their first SIM.')
journey_sheet('2_mnp',
              identity_note='No logged/guest split — same reason as New SIM.')
journey_sheet('3_recharge', identity_block('3_recharge'))
journey_sheet('4_invoice',  identity_block('4_invoice'))

# ⑤ Checkout — the one journey with a real logged/guest split
ws = journey_sheet('5_checkout',
                   identity_note='This journey DOES record logged vs guest directly '
                                 '(checkouts.checkoutable_type).')
r = ws.max_row + 3
ws.cell(row=r, column=1, value='Logged in vs guest (recorded, not inferred)').font = Font(bold=True, size=12, color=DARK)
r += 2
CO_COLS = ['checkout_user', 'checkout_guest', 'checkout_anonymoususer', 'checkout_deleteduser']
head(ws, r, ['Month'] + [c.replace('checkout_', '') for c in CO_COLS] + ['Total']); r += 1
cototal = collections.Counter()
for m in months:
    c = collections.Counter()
    for row in data[m]['journey']:
        if row['journey'] == '5_checkout':
            c[row['identity']] += row['n']; cototal[row['identity']] += row['n']
    body(ws, r, [m] + [c.get(k, 0) for k in CO_COLS] + [sum(c.get(k, 0) for k in CO_COLS)]); r += 1
body(ws, r, ['TOTAL'] + [cototal.get(k, 0) for k in CO_COLS] +
     [sum(cototal.get(k, 0) for k in CO_COLS)], bold=True)

# ---------------------------------------------------------------- 5 · failure reasons
ws = sheet('Failure reasons', [34, 44, 12, 14, 12, 14])
ws['A1'] = 'Why payments failed — 12 months, by journey'
ws['A1'].font = Font(bold=True, size=13, color=DARK)
ws['A2'] = '"(no message)" = the gateway returned no reason at all. Failed payments only.'
ws['A2'].font = Font(italic=True, size=10, color=GREY)
head(ws, 4, ['Journey', 'Reason (code · message)', 'Vendor', 'Rail', 'Payments', 'Value SAR'])
agg = collections.Counter(); sar = collections.Counter()
for m in months:
    for row in data[m]['reasons']:
        k = (row['journey'], row['reason'], row['vendor'], row['rail'])
        agg[k] += row['n']; sar[k] += num(row['sar'])
r = 5
for k, n in agg.most_common(600):
    body(ws, r, [JOURNEY_LABEL.get(k[0], k[0]), k[1], k[2], k[3], n, sar[k]], money=(6,)); r += 1

# reasons by month — the trend BI will want
ws = sheet('Reasons by month', [44] + [12] * (len(months) + 1))
ws['A1'] = 'Top 25 failure reasons, per month'
ws['A1'].font = Font(bold=True, size=13, color=DARK)
head(ws, 3, ['Reason'] + months + ['Total'])
bym = collections.defaultdict(collections.Counter); rt = collections.Counter()
for m in months:
    for row in data[m]['reasons']:
        bym[row['reason']][m] += row['n']; rt[row['reason']] += row['n']
r = 4
for reason, t in rt.most_common(25):
    body(ws, r, [reason] + [bym[reason].get(m, 0) for m in months] + [t]); r += 1

# ---------------------------------------------------------------- 6 · vendor / rail / platform
def cut_sheet(title, key, label):
    ws = sheet(title, [22] + [14] * 6)
    ws['A1'] = f'{label} × outcome — 12 months'
    ws['A1'].font = Font(bold=True, size=13, color=DARK)
    head(ws, 3, [label] + [o.replace('_', ' ').title() for o in OUTCOMES] + ['Total'])
    agg = collections.defaultdict(collections.Counter)
    for m in months:
        for row in data[m][key]:
            agg[row[label.lower().split()[0]]][row['outcome']] += row['n']
    r = 4
    for k in sorted(agg, key=lambda x: -sum(agg[x].values())):
        a = agg[k]
        body(ws, r, [k] + [a.get(o, 0) for o in OUTCOMES] + [sum(a.values())]); r += 1

cut_sheet('By vendor',   'byVendor',   'Vendor')
cut_sheet('By rail',     'byRail',     'Rail')
cut_sheet('By platform', 'byPlatform', 'Platform')

wb.save(OUT)
print(f'✓ {OUT}  ·  {len(wb.sheetnames)} sheets: {", ".join(wb.sheetnames)}')
