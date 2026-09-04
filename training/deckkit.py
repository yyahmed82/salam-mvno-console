"""
deckkit — shared drawing toolkit for the Salam Console workshop deck.

Light-technical theme + primitives for real diagrams (nodes, arrows, lanes, timelines,
state machines, matrices). Keeping this separate keeps the slide file readable.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
import os, glob

# ---------------- palette ----------------
BG      = RGBColor(0xFF, 0xFF, 0xFF)
INK     = RGBColor(0x0F, 0x17, 0x2A)
BODY    = RGBColor(0x33, 0x41, 0x55)
MUTED   = RGBColor(0x64, 0x74, 0x8B)
FAINT   = RGBColor(0x94, 0xA3, 0xB8)
GREEN   = RGBColor(0x00, 0x8A, 0x47)
GREENBG = RGBColor(0xEC, 0xFD, 0xF3)
DARK    = RGBColor(0x0A, 0x0A, 0x0F)
DARK2   = RGBColor(0x15, 0x17, 0x21)
CODEBG  = RGBColor(0x1E, 0x29, 0x3B)
CODEFG  = RGBColor(0xE2, 0xE8, 0xF0)
PANEL   = RGBColor(0xF8, 0xFA, 0xFC)
LINE    = RGBColor(0xE2, 0xE8, 0xF0)
AMBER   = RGBColor(0xB4, 0x53, 0x09)
AMBERBG = RGBColor(0xFF, 0xFB, 0xEB)
RED     = RGBColor(0xB9, 0x1C, 0x1C)
REDBG   = RGBColor(0xFE, 0xF2, 0xF2)
BLUE    = RGBColor(0x1D, 0x4E, 0xD8)
BLUEBG  = RGBColor(0xEF, 0xF6, 0xFF)
PURPLE  = RGBColor(0x6D, 0x28, 0xD9)
PURPBG  = RGBColor(0xF5, 0xF3, 0xFF)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
FONT    = "Aptos"
MONO    = "Consolas"

W, H = Inches(13.333), Inches(7.5)


# ---------------- basics ----------------
def slide(prs, fill=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill; f.solid(); f.fore_color.rgb = fill
    return s


def text(s, x, y, w, h, txt, size=12, color=BODY, bold=False, align=PP_ALIGN.LEFT,
         spacing=1.15, font=FONT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, ln in enumerate(str(txt).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        f = r.font; f.name = font; f.size = Pt(size); f.bold = bold
        f.color.rgb = color; f.italic = italic
    return tb


def box(s, x, y, w, h, fill=PANEL, line=LINE, radius=True, lw=0.75):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    try: sh.adjustments[0] = 0.06
    except Exception: pass
    return sh


def rule(s, x, y, w=0.7, color=GREEN, h=0.035):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()
    sh.shadow.inherit = False; return sh


def code(s, x, y, w, h, lines, size=10, fg=CODEFG, bg=CODEBG):
    box(s, x, y, w, h, bg, None)
    text(s, x + 0.22, y + 0.16, w - 0.44, h - 0.32, lines, size, fg, font=MONO, spacing=1.3)


def callout(s, x, y, w, h, label, body, kind="key", size=11):
    pal = {"key": (GREENBG, GREEN), "warn": (AMBERBG, AMBER), "stop": (REDBG, RED),
           "info": (BLUEBG, BLUE), "note": (PURPBG, PURPLE)}[kind]
    box(s, x, y, w, h, pal[0], pal[1])
    text(s, x + 0.26, y + 0.16, w - 0.52, 0.24, label.upper(), 9.5, pal[1], bold=True)
    text(s, x + 0.26, y + 0.46, w - 0.52, h - 0.62, body, size, INK, spacing=1.22)


# ---------------- diagram primitives ----------------
def node(s, x, y, w, h, title, sub=None, fill=WHITE, edge=BLUE, tsize=11, ssize=8.5,
         tcolor=None, radius=True):
    """A labelled diagram node."""
    box(s, x, y, w, h, fill, edge, radius, lw=1.1)
    if sub:
        text(s, x, y + h / 2 - 0.28, w, 0.26, title, tsize, tcolor or INK, bold=True, align=PP_ALIGN.CENTER)
        text(s, x, y + h / 2 + 0.01, w, 0.24, sub, ssize, MUTED, align=PP_ALIGN.CENTER)
    else:
        text(s, x, y + h / 2 - 0.13, w, 0.26, title, tsize, tcolor or INK, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return (x, y, w, h)


def arrow(s, x1, y1, x2, y2, color=FAINT, width=1.6, dashed=False, head=True):
    """Straight connector with an arrowhead."""
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        from pptx.oxml.ns import qn
        from lxml import etree
        tail = etree.SubElement(ln, qn('a:tailEnd'))
        tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    return c


def arrow_r(s, frm, to, color=FAINT, width=1.6, dashed=False, gap=0.06):
    """Arrow from right edge of node `frm` to left edge of node `to` (nodes are (x,y,w,h))."""
    x1 = frm[0] + frm[2] + gap; y1 = frm[1] + frm[3] / 2
    x2 = to[0] - gap;           y2 = to[1] + to[3] / 2
    return arrow(s, x1, y1, x2, y2, color, width, dashed)


def arrow_d(s, frm, to, color=FAINT, width=1.6, dashed=False, gap=0.06):
    """Arrow from bottom edge of `frm` to top edge of `to`."""
    x1 = frm[0] + frm[2] / 2; y1 = frm[1] + frm[3] + gap
    x2 = to[0] + to[2] / 2;   y2 = to[1] - gap
    return arrow(s, x1, y1, x2, y2, color, width, dashed)


def edge_label(s, x, y, w, txt, color=MUTED, size=8.5, bg=None):
    if bg: box(s, x, y - 0.02, w, 0.24, bg, None)
    text(s, x, y, w, 0.22, txt, size, color, align=PP_ALIGN.CENTER)


def lane(s, x, y, w, h, label, fill=PANEL, edge=LINE, lcolor=MUTED, label_above=False):
    """Horizontal swimlane. label_above=True puts the caption OUTSIDE the band — use it whenever
    content starts near the left edge, otherwise a long caption wraps under the first node."""
    box(s, x, y, w, h, fill, edge)
    if label_above:
        text(s, x, y - 0.26, w, 0.22, label.upper(), 9, lcolor, bold=True)
    else:
        text(s, x + 0.16, y + 0.12, 1.4, 0.24, label.upper(), 9, lcolor, bold=True)
    return (x, y, w, h)


def timeline(s, x, y, w, points, color=BLUE, dot=0.14, above=True):
    """points = [(frac 0..1, 'label', 'sub')]  — a horizontal timeline."""
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.022))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background(); ln.shadow.inherit = False
    for frac, label, sub in points:
        cx = x + w * frac
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - dot / 2), Inches(y - dot / 2 + 0.011),
                               Inches(dot), Inches(dot))
        d.fill.solid(); d.fill.fore_color.rgb = color; d.line.color.rgb = WHITE
        d.line.width = Pt(1.5); d.shadow.inherit = False
        if above:
            text(s, cx - 1.0, y - 0.62, 2.0, 0.24, label, 10, INK, bold=True, align=PP_ALIGN.CENTER)
            text(s, cx - 1.1, y - 0.36, 2.2, 0.22, sub, 8.5, MUTED, align=PP_ALIGN.CENTER)
        else:
            text(s, cx - 1.0, y + 0.22, 2.0, 0.24, label, 10, INK, bold=True, align=PP_ALIGN.CENTER)
            text(s, cx - 1.1, y + 0.48, 2.2, 0.22, sub, 8.5, MUTED, align=PP_ALIGN.CENTER)


def table(s, x, y, w, rows, colw, size=10, header=True, rowh=0.34, zebra=True,
          hfill=RGBColor(0xF1, 0xF5, 0xF9), colors=None):
    """rows = list of lists. colors = optional dict {(r,c): RGBColor} for cell text."""
    for r, row in enumerate(rows):
        yy = y + r * rowh
        if r == 0 and header: box(s, x, yy, w, rowh, hfill, LINE)
        elif zebra and r % 2 == 0: box(s, x, yy, w, rowh, RGBColor(0xFB, 0xFC, 0xFD), None)
        cx = x
        for c, cell in enumerate(row):
            col = (colors or {}).get((r, c), INK if r == 0 else BODY)
            text(s, cx + 0.13, yy + 0.06, colw[c] - 0.18, rowh - 0.08, str(cell), size, col,
                 bold=(r == 0 and header))
            cx += colw[c]
    return y + len(rows) * rowh


def kpi(s, x, y, w, h, value, label, color=GREEN, vsize=26, lsize=9):
    box(s, x, y, w, h)
    text(s, x, y + 0.16, w, 0.5, value, vsize, color, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, y + h - 0.34, w, 0.24, label.upper(), lsize, MUTED, align=PP_ALIGN.CENTER)


def shot(s, slot, caption, shots_dir, x=6.95, y=2.05, w=5.75, h=3.9):
    hit = [f for f in glob.glob(os.path.join(shots_dir, slot + ".*"))
           if f.lower().endswith((".png", ".jpg", ".jpeg"))]
    if hit:
        pic = s.shapes.add_picture(hit[0], Inches(x), Inches(y), width=Inches(w))
        if pic.height > Inches(h):
            el = pic._element; el.getparent().remove(el)
            pic = s.shapes.add_picture(hit[0], Inches(x), Inches(y), height=Inches(h))
            pic.left = Inches(x + (w - pic.width.inches) / 2)
        return pic
    box(s, x, y, w, h, PANEL, LINE)
    text(s, x, y + h / 2 - 0.35, w, 0.3, "▢", 18, RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.CENTER)
    text(s, x, y + h / 2 - 0.02, w, 0.3, slot + ".png", 11, MUTED, align=PP_ALIGN.CENTER)
    text(s, x, y + h / 2 + 0.3, w, 0.5, caption, 9.5, MUTED, align=PP_ALIGN.CENTER)


def qrpic(s, x, y, size, data, color="008A47"):
    """QR code as a picture. Silently no-ops if segno isn't installed or `data` is a placeholder,
    so the deck always builds — the URL text beside it is the real content."""
    if not data or str(data).startswith("PASTE_"):
        return False
    try:
        import segno, io
        buf = io.BytesIO()
        segno.make(str(data), error="m").save(buf, kind="png", scale=8, border=0,
                                              dark="#" + color, light="#FFFFFF")
        buf.seek(0)
        s.shapes.add_picture(buf, Inches(x), Inches(y), Inches(size), Inches(size))
        return True
    except Exception:
        return False


def linkrow(s, x, y, w, h, label, url, sub, cfgok=True, accent=BLUE, accentbg=BLUEBG, qr=True):
    """A resource row: label + live URL + optional QR. Renders a 'to be confirmed' state
    when the config still holds a placeholder, so nothing looks broken in a dry run."""
    ok = bool(url) and not str(url).startswith("PASTE_")
    box(s, x, y, w, h, accentbg if ok else PANEL, accent if ok else LINE)
    q = qr and qrpic(s, x + w - h + 0.14, y + 0.14, h - 0.28, url, str(accent))
    tw = w - (h + 0.3 if q else 0.5)
    text(s, x + 0.26, y + 0.16, tw, 0.24, label.upper(), 9.5, accent if ok else MUTED, bold=True)
    text(s, x + 0.26, y + 0.44, tw, 0.3, url if ok else "— link to be added before session 1 —",
         10.5 if ok else 10.5, INK if ok else FAINT, font=MONO if ok else FONT,
         italic=not ok)
    text(s, x + 0.26, y + 0.78, tw, 0.34, sub, 9.5, MUTED, spacing=1.15)


def banner(s, x, y, w, txt, kind="key", size=12, h=0.5):
    """Single-line emphasis. callout() reserves ~0.62in for its label row, so using it for a
    one-liner leaves the text outside the box — this is the right shape for that case."""
    pal = {"key": (GREENBG, GREEN), "warn": (AMBERBG, AMBER), "stop": (REDBG, RED),
           "info": (BLUEBG, BLUE), "note": (PURPBG, PURPLE)}[kind]
    box(s, x, y, w, h, pal[0], pal[1])
    text(s, x + 0.3, y, w - 0.6, h, txt, size, INK, anchor=MSO_ANCHOR.MIDDLE)
    return y + h


def checklist(s, x, y, w, items, title=None, colr=GREEN, bgc=GREENBG, rowh=0.42, size=11.5):
    """Sign-off list with empty tick boxes — the artefact L2 actually fills in."""
    top = y
    if title:
        text(s, x, y, w, 0.26, title.upper(), 9.5, colr, bold=True)
        top = y + 0.34
    for i, it in enumerate(items):
        yy = top + i * rowh
        bx = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(yy + 0.04),
                                Inches(0.24), Inches(0.24))
        bx.fill.solid(); bx.fill.fore_color.rgb = WHITE
        bx.line.color.rgb = colr; bx.line.width = Pt(1.1); bx.shadow.inherit = False
        try: bx.adjustments[0] = 0.15
        except Exception: pass
        text(s, x + 0.36, yy, w - 0.4, rowh - 0.05, it, size, INK, spacing=1.15)
    return top + len(items) * rowh


def verdict(s, x, y, w, h, label, question, options_, qshare=0.52):
    """A decision the room must make before moving on — the point of a review session.

    HORIZONTAL by design: question on the left, option chips on the right. A stacked layout needs
    ~1.2in and these blocks usually sit in whatever strip is left at the bottom of a slide, so
    stacking silently drew the options over the question. h is clamped to a workable minimum.
    """
    h = max(h, 0.72)
    box(s, x, y, w, h, AMBERBG, AMBER)
    qw = w * qshare - 0.4
    text(s, x + 0.28, y + 0.13, qw, 0.22, label.upper(), 9, AMBER, bold=True)
    text(s, x + 0.28, y + 0.36, qw, h - 0.48, question, 12.5, INK, bold=True, spacing=1.15)
    n = len(options_)
    ox0 = x + w * qshare
    ow = (w - w * qshare - 0.28 - 0.14 * (n - 1)) / n
    ch = min(0.44, h - 0.24)
    for i, o in enumerate(options_):
        ox = ox0 + i * (ow + 0.14)
        box(s, ox, y + (h - ch) / 2, ow, ch, WHITE, AMBER)
        text(s, ox + 0.06, y + (h - ch) / 2, ow - 0.12, ch, o, 10, AMBER, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ---------------- slide chrome ----------------
def head(s, module, title, sub=None, subw=12.1, badge=None):
    text(s, 0.62, 0.4, 8.0, 0.2, module.upper(), 9.5, GREEN, bold=True)
    text(s, 0.62, 0.68, 11.0, 0.45, title, 24, INK, bold=True)
    rule(s, 0.62, 1.24)
    if sub: text(s, 0.62, 1.44, subw, 0.5, sub, 12, MUTED)
    if badge:
        box(s, 11.35, 0.38, 1.37, 0.36, GREENBG, GREEN)
        text(s, 11.35, 0.46, 1.37, 0.22, badge.upper(), 9, GREEN, bold=True, align=PP_ALIGN.CENTER)


def section(s, num, title, sub, modules):
    """Dark session divider listing the modules it contains."""
    text(s, 0.62, 1.5, 6.0, 0.35, num.upper(), 12, GREEN, bold=True)
    text(s, 0.62, 1.9, 11.5, 0.8, title, 34, WHITE, bold=True)
    rule(s, 0.62, 2.95, 1.1, GREEN)
    text(s, 0.62, 3.2, 10.5, 0.45, sub, 14, RGBColor(0xE2, 0xE8, 0xF0))
    # spacing adapts to the module count so a 6-row session can't run off the slide
    n = len(modules)
    step = 0.62 if n <= 5 else 0.56
    top = 4.0 if n <= 5 else 3.78
    for i, (code_, name, mins) in enumerate(modules):
        y = top + i * step
        box(s, 0.62, y, 11.5, 0.48 if n > 5 else 0.52, DARK2, RGBColor(0x27, 0x2B, 0x38))
        text(s, 0.88, y + 0.14, 0.9, 0.24, code_, 10.5, GREEN, bold=True)
        text(s, 1.85, y + 0.14, 8.4, 0.24, name, 11.5, RGBColor(0xE2, 0xE8, 0xF0))
        text(s, 10.5, y + 0.14, 1.4, 0.24, mins, 10, RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.RIGHT)


def labhead(s, num, title, mins, color=GREEN):
    box(s, 0, 0, 13.333, 1.05, color, None, radius=False)
    text(s, 0.62, 0.22, 9.0, 0.26, f"LAB {num}  ·  {mins}  ·  you drive", 10.5,
         RGBColor(0xD1, 0xFA, 0xE5), bold=True)
    text(s, 0.62, 0.52, 11.5, 0.4, title, 19, WHITE, bold=True)


def pollhead(s, n, title, mins):
    box(s, 0, 0, 13.333, 1.05, BLUE, None, radius=False)
    text(s, 0.62, 0.22, 9.0, 0.26, f"LIVE POLL {n}  ·  {mins}  ·  everyone answers", 10.5,
         RGBColor(0xDB, 0xEA, 0xFE), bold=True)
    text(s, 0.62, 0.52, 11.5, 0.4, title, 19, WHITE, bold=True)


def options(s, opts, x=0.62, y=2.6, w=5.9, step=0.72, letters="ABCD"):
    for i, o in enumerate(opts):
        yy = y + i * step
        box(s, x, yy, w, 0.6)
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(yy + 0.14),
                               Inches(0.32), Inches(0.32))
        d.fill.solid(); d.fill.fore_color.rgb = BLUEBG; d.line.color.rgb = BLUE
        d.shadow.inherit = False
        text(s, x + 0.2, yy + 0.19, 0.32, 0.24, letters[i], 10.5, BLUE, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.66, yy + 0.17, w - 0.9, 0.3, o, 12, INK)


def reveal(s, answer, why, extra=None):
    box(s, 6.85, 2.6, 5.85, 0.6, GREENBG, GREEN)
    text(s, 7.1, 2.75, 5.4, 0.3, f"✓  ANSWER: {answer}", 13, GREEN, bold=True)
    text(s, 6.85, 3.4, 5.85, 0.25, "WHY IT MATTERS", 9.5, MUTED, bold=True)
    text(s, 6.85, 3.68, 5.85, 1.6, why, 11.5, BODY, spacing=1.25)
    if extra: callout(s, 6.85, 5.4, 5.85, 1.3, "remember", extra, "key", 11)
