#!/usr/bin/env python3
"""
Builds a Kahoot-compatible quiz spreadsheet (also usable as a source for Slido / Mentimeter /
Microsoft Forms) so the in-deck polls can be run as a scored, competitive game.

Kahoot's bulk-import template columns:
    Question | Answer 1 | Answer 2 | Answer 3 | Answer 4 | Time limit (sec) | Correct answer(s)

Rules Kahoot enforces on import:
  · question ≤ 120 chars   · each answer ≤ 75 chars   · time limit from {5,10,20,30,60,90,120,240}
  · correct answer(s) as the answer NUMBER (1-4), comma separated for multi-correct
This script validates all of that before writing, so the import can't be rejected.

USAGE   python3 build_quiz_import.py
OUTPUT  Salam-Console-Quiz-Kahoot.xlsx   (4 rounds, 24 questions)
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Salam-Console-Quiz-Kahoot.xlsx")

# (question, [a1..a4], correct_index_1based, seconds, round)
Q = [
# ── ROUND 1 · Foundations ────────────────────────────────────────────────────
("Where does the console read its data from?",
 ["Production directly", "A replica synced from prod", "A nightly CSV export", "The BSS directly"], 2, 20, 1),
("Can the console write to production?",
 ["Yes, with Super Admin", "Yes, it can refund payments", "No — read-only by construction", "Only activations"], 3, 20, 1),
("How fresh is the data normally?",
 ["Real time", "About 5 minutes", "One hour", "Yesterday"], 2, 20, 1),
("Who can unmask customer PII?",
 ["Anyone logged in", "Any L2", "Super Admin and L3 Digital only", "Nobody"], 3, 20, 1),
("What happens when someone unmasks PII?",
 ["Nothing special", "Written to the audit log with their name", "The customer is notified", "It is cached"], 2, 20, 1),
("A red staleness banner appears. What do you do?",
 ["Ignore it, numbers look fine", "Refresh until it clears", "Treat figures as NOT live and say so", "Restart the console"], 3, 30, 1),
# ── ROUND 2 · Troubleshooting ────────────────────────────────────────────────
("Payment pending 45 min WITH a gateway commit response. What is it?",
 ["Abandoned by the customer", "Stuck — may be charged, escalate", "A duplicate charge", "Completely normal"], 2, 30, 2),
("Payment pending with NO commit response. What is it?",
 ["Stuck — escalate now", "The customer never finished", "A gateway outage", "A refund in progress"], 2, 30, 2),
("Why does the duplicate tile say SUSPECTED?",
 ["It is a beta feature", "Confirmation needs gateway reconciliation", "It is usually wrong", "Legal wording"], 2, 20, 2),
("Blended success is 84% but a gateway is dead. How?",
 ["The dashboard is broken", "Traffic failed over to another gateway", "84% is actually bad", "That gateway is excluded"], 2, 30, 2),
("Which of these is an INCIDENT, not a single customer case?",
 ["One card declined", "One abandoned checkout", "20 customers failing the same step in 10 min", "A forgotten password"], 3, 30, 2),
("BSS write path is green but customers say plans won't load. Why?",
 ["They are wrong", "READ-path SOAP faults, logged in another database", "Their app is old", "The SIM is faulty"], 2, 30, 2),
# ── ROUND 3 · Alerting ───────────────────────────────────────────────────────
("Semati fails about 55% on a normal day. When should an alert fire?",
 ["On any failure", "Above 10%", "Meaningfully above the 55% baseline", "Never"], 3, 30, 3),
("What does acknowledging an incident actually do?",
 ["Marks it resolved", "Stops the escalation ladder climbing", "Emails the customer", "Nothing, cosmetic"], 2, 20, 3),
("Why does correlation/suppression exist?",
 ["To store fewer alerts", "So one outage pages once, not ten times", "To hide alerts", "To save disk"], 2, 20, 3),
("Where does the runbook for an incident live?",
 ["In a wiki", "In the WhatsApp group", "Attached to the alert itself", "In someone's head"], 3, 20, 3),
("An SLO shows AT RISK. What does that mean?",
 ["Already breached", "Still met, but error budget nearly used", "No data", "Target is wrong"], 2, 30, 3),
("What does the anomaly detector compare today against?",
 ["A fixed threshold", "The same hour of the week, recent weeks", "Last year", "The team average"], 2, 30, 3),
# ── ROUND 4 · Depth ──────────────────────────────────────────────────────────
("Reseller channel numbers must be attributed from…",
 ["plan_channels (which channels MAY sell)", "The order's own recorded origin", "UTM only", "Manual reporting"], 2, 30, 4),
("Why did plan_channels give wrong numbers?",
 ["It was not synced", "One order counted once per enabled channel", "It had no index", "It was too slow"], 2, 30, 4),
("What should you do BEFORE changing an alert threshold?",
 ["Guess a round number", "Look at the metric chart and its band", "Disable the rule first", "Ask the vendor"], 2, 20, 4),
("An alert cries wolf nightly at 03:00. Best fix?",
 ["Delete the rule", "Disable it forever", "Raise the min-sample gate", "Ignore it"], 3, 30, 4),
("Why must everything run in UTC?",
 ["It does not matter", "Naive UTC columns shift if the session is not UTC", "For log formatting", "Legal requirement"], 2, 30, 4),
("Where does your data go when you ask Yusr a question?",
 ["To OpenAI", "Nowhere — the model runs on-premise", "To Teams", "To the vendor"], 2, 20, 4),
]

VALID_TIMES = {5, 10, 20, 30, 60, 90, 120, 240}
ROUND_NAMES = {1: "Foundations", 2: "Troubleshooting", 3: "Alerting", 4: "Depth"}

# ---- validate against Kahoot's import rules before writing ----
errs = []
for i, (q, ans, correct, secs, rnd) in enumerate(Q, 1):
    if len(q) > 120: errs.append(f"Q{i}: question {len(q)} chars (max 120)")
    if len(ans) != 4: errs.append(f"Q{i}: needs exactly 4 answers")
    for j, a in enumerate(ans, 1):
        if len(a) > 75: errs.append(f"Q{i} A{j}: {len(a)} chars (max 75)")
    if not 1 <= correct <= 4: errs.append(f"Q{i}: correct index {correct}")
    if secs not in VALID_TIMES: errs.append(f"Q{i}: time {secs} not allowed")
if errs:
    print("✗ validation failed:"); [print("   ", e) for e in errs]; raise SystemExit(1)

wb = Workbook()

# one sheet per round → import each as a separate Kahoot game, or combine
for rnd in sorted(ROUND_NAMES):
    ws = wb.active if rnd == 1 else wb.create_sheet()
    ws.title = f"Round {rnd} · {ROUND_NAMES[rnd]}"[:31]
    hdr = ["Question", "Answer 1", "Answer 2", "Answer 3", "Answer 4",
           "Time limit (sec)", "Correct answer(s)"]
    ws.append(hdr)
    for c in range(1, 8):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="008A47")
        cell.alignment = Alignment(vertical="center")
    for q, ans, correct, secs, r in Q:
        if r != rnd: continue
        ws.append([q, *ans, secs, correct])
    widths = [62, 34, 34, 34, 34, 16, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

wb.save(OUT)
per = {r: sum(1 for *_x, rr in Q if rr == r) for r in ROUND_NAMES}
print(f"✓ {OUT}")
print("  " + " · ".join(f"Round {r} ({ROUND_NAMES[r]}): {per[r]} Q" for r in sorted(ROUND_NAMES)))
print(f"  total: {len(Q)} questions — all validated against Kahoot import rules")
