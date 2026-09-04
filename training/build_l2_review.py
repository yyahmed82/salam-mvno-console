#!/usr/bin/env python3
"""
Salam Digital Console — L2 REVIEW SERIES (3 sessions, this week).

DIFFERENT PURPOSE FROM THE L1 WORKSHOP. L2 are not the audience here — they are the reviewers.
Each session ends in decisions L2 sign off: are the numbers right, are the thresholds right, is the
console actually useful in an incident. Between sessions they use it for real and log findings.

USAGE   python3 build_l2_review.py
OUTPUT  Salam-Console-L2-Review.pptx
"""
import os, glob, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from deckkit import *

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "shots")
OUT = os.path.join(HERE, "Salam-Console-L2-Review.pptx")
CFG = json.load(open(os.path.join(HERE, "workshop.config.json"), encoding="utf-8"))
TEAMS = CFG["teams"]

prs = Presentation(); prs.slide_width, prs.slide_height = W, H
S = lambda fill=BG: slide(prs, fill)

# live system facts — keep in step with seedRules.js / metrics.js
N_RULES, N_METRICS, N_P1, N_P2, N_P3 = 43, 33, 15, 24, 4


# ══════════════════════════════════ TITLE ══════════════════════════════════
s = S(DARK)
text(s, 0, 1.95, 13.333, 0.25, "L2 REVIEW SERIES  ·  3 SESSIONS  ·  THIS WEEK", 11, GREEN, bold=True, align=PP_ALIGN.CENTER)
text(s, 0, 2.35, 13.333, 0.95, "Salam Digital Console", 42, WHITE, bold=True, align=PP_ALIGN.CENTER)
rule(s, 5.9, 3.45, 1.55, GREEN)
text(s, 0, 3.72, 13.333, 0.45, "You are not the audience. You are the reviewers.", 17, RGBColor(0xE2,0xE8,0xF0), align=PP_ALIGN.CENTER)
text(s, 0, 4.3, 13.333, 0.6, "Three sessions to decide whether the data is true, the alerts are right,\nand the console is worth opening during an incident.",
     12.5, RGBColor(0x94,0xA3,0xB8), align=PP_ALIGN.CENTER, spacing=1.35)
for i, (v, l) in enumerate([("3", "sessions"), (str(N_RULES), "rules to review"), (str(N_METRICS), "metrics"),
                            ("5", "days of live use")]):
    x = 2.63 + i * 2.15
    box(s, x, 5.5, 1.95, 1.05, DARK2, RGBColor(0x27,0x2B,0x38))
    text(s, x, 5.69, 1.95, 0.4, v, 20, GREEN, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, 6.15, 1.95, 0.22, l.upper(), 8.5, RGBColor(0x94,0xA3,0xB8), align=PP_ALIGN.CENTER)

# ══════════════════════════════════ WHY L2 FIRST ══════════════════════════════════
s = S()
head(s, "why you, why now", "L2 before L1 — deliberately",
     "We could train L1 first and get more people using it sooner. That would be the wrong order.")
for i, (t, b, c, bgc) in enumerate([
    ("L1 will trust what the console says",
     "If a number is wrong, L1 has no way to know. They will quote it to a customer, or close a case on it. "
     "Whatever we ship to L1 has to be right BEFORE they see it — and you are the only people who can tell.",
     RED, REDBG),
    ("You already know what the truth looks like",
     "You have watched Semati fail at 55% on a normal day. You know activation errors that mean nothing and "
     "ones that mean everything. That judgement is not in the code — it is in your heads, and I need it in the rules.",
     GREEN, GREENBG),
    ("A week of real use beats any amount of demo",
     "Between sessions you use it on real cases. Everything that looks wrong, feels noisy, or is missing gets "
     "logged. We review it together the next morning. That loop is the actual deliverable.",
     BLUE, BLUEBG)]):
    y = 2.15 + i * 1.55
    box(s, 0.62, y, 12.1, 1.4, bgc, c)
    text(s, 0.92, y + 0.2, 11.5, 0.28, t, 14, INK, bold=True)
    text(s, 0.92, y + 0.56, 11.5, 0.7, b, 11.5, BODY, spacing=1.25)
banner(s, 0.62, 6.9, 12.1, "L1 workshops start only once this review closes. If that means a delay, the delay is the right call.", "key", 12)

# ══════════════════════════════════ THE WEEK ══════════════════════════════════
s = S()
head(s, "the week", "Three sessions, five days, one feedback loop",
     "The work between the sessions matters more than the sessions.")
cols = [
    ("SESSION 1", "Is the data TRUE?", GREEN, GREENBG,
     ["Architecture & lineage", "Sync, lag & staleness", "Every definition on the dashboard",
      "Attribution traps", "What we can and cannot see"],
     "Sign off 12 definitions"),
    ("SESSION 2", "Are the ALERTS right?", BLUE, BLUEBG,
     ["How a rule evaluates", f"All {N_RULES} rules by severity", "Baselines vs absolute thresholds",
      "INC0016809 — the miss", "Correlation, escalation, SLA"],
     "Sign off every threshold"),
    ("SESSION 3", "Is it USEFUL?", PURPLE, PURPBG,
     ["Troubleshoot board", "Timeline & Subscriber 360", "Three real incidents, replayed",
      "Yusr assistant", "Ownership & cadence"],
     "Verdict + ranked backlog"),
]
for i, (num, title, c, bgc, items, out) in enumerate(cols):
    x = 0.62 + i * 4.1
    box(s, x, 2.15, 3.9, 3.7, bgc, c)
    text(s, x + 0.28, 2.35, 3.3, 0.24, num, 10, c, bold=True)
    text(s, x + 0.28, 2.65, 3.4, 0.34, title, 15, INK, bold=True)
    rule(s, x + 0.28, 3.12, 0.45, c, 0.03)
    for j, it in enumerate(items):
        text(s, x + 0.28, 3.34 + j * 0.42, 3.4, 0.38, "· " + it, 10.5, BODY, spacing=1.15)
    box(s, x + 0.28, 5.5, 3.34, 0.2, bgc, None)
    text(s, x + 0.28, 5.46, 3.4, 0.24, "→ " + out, 10.5, c, bold=True)
box(s, 0.62, 6.05, 12.1, 1.15, PANEL, LINE)
text(s, 0.92, 6.24, 11.6, 0.26, "BETWEEN SESSIONS — the part that decides whether this works", 10, GREEN, bold=True)
text(s, 0.92, 6.56, 11.6, 0.55,
     "Use the console on real cases · log anything wrong, noisy or missing in the review log · "
     "15-minute stand-up each morning to triage what you logged. Nothing is too small — the reseller "
     "numbers were wrong for weeks because the join looked fine.", 11, BODY, spacing=1.25)

# ══════════════════════════════════ HOW WE EXCHANGE ══════════════════════════════════
s = S()
head(s, "the loop", "How we exchange findings",
     "One place for everything, so no finding is lost in a DM.")
lane(s, 0.62, 2.35, 12.1, 1.25, "you, during the day", label_above=True)
n1 = node(s, 1.05, 2.6, 2.5, 0.75, "Use it on a real case", "instead of the old way", WHITE, GREEN)
n2 = node(s, 4.3, 2.6, 2.5, 0.75, "Something looks off", "number · alert · gap", WHITE, AMBER)
n3 = node(s, 7.55, 2.6, 2.5, 0.75, "Log it", "review log, 30 seconds", WHITE, BLUE)
n4 = node(s, 10.3, 2.6, 2.1, 0.75, "Keep working", None, WHITE, FAINT)
for a, b in ((n1, n2), (n2, n3), (n3, n4)): arrow_r(s, a, b)

lane(s, 0.62, 4.05, 12.1, 1.25, "us, next morning", label_above=True)
m1 = node(s, 1.05, 4.3, 2.5, 0.75, "Triage together", "real / not real", WHITE, GREEN)
m2 = node(s, 4.3, 4.3, 2.5, 0.75, "Classify", "data · rule · UX · gap", WHITE, BLUE)
m3 = node(s, 7.55, 4.3, 2.5, 0.75, "Fix or schedule", "same day if data is wrong", WHITE, PURPLE)
m4 = node(s, 10.3, 4.3, 2.1, 0.75, "Close the loop", "you see the fix", WHITE, GREEN)
for a, b in ((m1, m2), (m2, m3), (m3, m4)): arrow_r(s, a, b)

table(s, 0.62, 5.45, 12.1,
      [["Finding type", "Example", "Who fixes", "Target"],
       ["Data wrong", "a KPI disagrees with BSS / prod", "console owner", "same day — this is the worst kind"],
       ["Alert noisy", "fires nightly, never actionable", "L2 + owner, together", "next session"],
       ["Alert missing", "an incident nothing caught", "console owner", "next session"],
       ["Hard to use", "took 6 clicks to answer one question", "console owner", "backlog, ranked in S3"]],
      [2.3, 4.3, 2.3, 3.2], 10.5)
banner(s, 0.62, 7.02, 12.1, "Ask in the channel, not in a DM — the channel is the record.", "key", 11.5, 0.44)

# ══════════════════════════════════ SESSION 1 DIVIDER ══════════════════════════════════
s = S(DARK)
section(s, "session 1 · day 1", "Is the data true?",
        "Every number on the dashboard, traced back to the row it came from",
        [("1.1", "Architecture & data lineage — where each number comes from", ""),
         ("1.2", "Sync, lag, and how you know the data is stale", ""),
         ("1.3", "Definitions review — the order flow, one box at a time", ""),
         ("1.4", "Attribution: the trap that made an unlaunched channel look live", ""),
         ("1.5", "Coverage — what the console can and cannot see", ""),
         ("", "REVIEW BLOCK — sign off 12 definitions", "")])

# ---- 1.1 lineage ----
s = S()
head(s, "1.1 · architecture", "Where every number comes from",
     "Four hops. If you do not trust a number, we walk it back along this line together.")
a = node(s, 0.62, 2.5, 2.3, 0.95, "Production", "selfcare Postgres", WHITE, FAINT)
b = node(s, 3.35, 2.5, 2.3, 0.95, "Replica", "prod-sync, ~5 min", WHITE, BLUE)
c_ = node(s, 6.08, 2.5, 2.3, 0.95, "Metrics", "33 registered", WHITE, PURPLE)
d = node(s, 8.81, 2.5, 2.3, 0.95, "Rules", f"{N_RULES} evaluated", WHITE, AMBER)
e = node(s, 11.0, 4.3, 1.72, 0.95, "You", None, WHITE, GREEN)
for x_, y_ in ((a, b), (b, c_), (c_, d)): arrow_r(s, x_, y_)
arrow(s, 9.96, 3.45, 11.5, 4.28, FAINT)
box(s, 0.62, 4.3, 9.9, 0.95, PANEL, LINE)
text(s, 0.9, 4.45, 9.4, 0.24, "READ-ONLY BY CONSTRUCTION", 9.5, GREEN, bold=True)
text(s, 0.9, 4.73, 9.4, 0.42, "The console has no write path to production. It cannot refund, activate, cancel or edit "
     "anything. Worst case it is WRONG — never destructive.", 11, BODY, spacing=1.2)
table(s, 0.62, 5.55, 12.1,
      [["Hop", "What can go wrong", "How you would notice"],
       ["prod → replica", "sync stalls, table wedged", "staleness banner + freshness check"],
       ["replica → metric", "SQL says something other than you assume", "the definitions review, today"],
       ["metric → rule", "threshold wrong for the baseline", "session 2"],
       ["rule → you", "fires but nobody acts", "session 3"]],
      [3.0, 5.2, 3.9], 10.5)

# ---- 1.2 sync & staleness ----
s = S()
head(s, "1.2 · freshness", "How you know the data is stale",
     "A wrong number is bad. A stale number that looks live is worse.")
timeline(s, 1.2, 2.9, 11.0, [
    (0.0, "prod write", "customer acts"),
    (0.3, "prod-sync", "every 5 min"),
    (0.58, "replica", "watermark advances"),
    (0.85, "console", "cached ≤120s")], BLUE)
for i, (t, b, k) in enumerate([
    ("Normal", "newest row under ~10 minutes old. The header shows 'updated Ns ago'.", "key"),
    ("Degraded", "10–45 min. Amber banner. Numbers are directionally right, not live — say so on a bridge.", "warn"),
    ("Stale", "over 45 min. Red banner. Treat every figure as NOT current, and do not quote it.", "stop")]):
    x = 0.62 + i * 4.1
    callout(s, x, 3.85, 3.9, 1.3, t, b, k, 10.5)
box(s, 0.62, 5.4, 12.1, 0.95, PANEL, LINE)
text(s, 0.9, 5.56, 11.6, 0.24, "THE WATCHDOG — why you should not have to think about this", 9.5, GREEN, bold=True)
text(s, 0.9, 5.86, 11.6, 0.42, "A cron job measures real lag every 5 min and only syncs when the replica is actually behind — "
     "retries with backoff, and re-pulls an explicit window if a watermark wedges. Healthy system = zero extra load on prod.",
     11, BODY, spacing=1.2)
verdict(s, 0.62, 6.5, 12.1, 0.9, "review", "Is a 10-minute lag acceptable for how you work?",
        ["Yes — 10/45 is right", "No — tighten", "Depends on the page"])

# ---- 1.3 definitions ----
s = S()
head(s, "1.3 · definitions", "The order flow, one box at a time",
     "This is the slowest part of the week and the most valuable. Every box is a definition you can veto.")
lane(s, 0.62, 2.3, 12.1, 1.15, "the lane you already know", label_above=True)
o = node(s, 1.0, 2.55, 1.75, 0.7, "Orders", None, WHITE, FAINT)
p1 = node(s, 3.05, 2.55, 1.75, 0.7, "Eligibility", None, WHITE, FAINT)
p2 = node(s, 5.1, 2.55, 1.75, 0.7, "Payment", None, WHITE, FAINT)
p3 = node(s, 7.15, 2.55, 1.75, 0.7, "SIM", None, WHITE, FAINT)
p4 = node(s, 9.2, 2.55, 1.75, 0.7, "Delivery", None, WHITE, FAINT)
p5 = node(s, 11.0, 2.55, 1.5, 0.7, "Activated", None, WHITE, GREEN)
for x_, y_ in ((o, p1), (p1, p2), (p2, p3), (p3, p4), (p4, p5)): arrow_r(s, x_, y_, gap=0.05)
table(s, 0.62, 3.6, 12.1,
      [["Box", "What it counts TODAY", "The question for you"],
       ["Eligibility pass", "is_eligible = true", "does this match how CITC/Semati decides?"],
       ["Total payment", "eligible AND a payment row exists", "should abandoned checkouts count here?"],
       ["Pending payment", "status pending/initiated", "how long before pending means STUCK?"],
       ["Physical SIM", "paid AND sim_type ≠ eSIM", "any other SIM types we are missing?"],
       ["Not assigned", "no delivery request AND not a partner", "is the partner split right? (new this week)"],
       ["Partner (tygo/soob)", "order origin = tygo or soob", "any other partner to add?"],
       ["Delivered", "delivery_state in the completed list", "is that list complete and current?"],
       ["Activated", "onboarding_orders.activated", "does BSS agree with this flag?"]],
      [2.5, 5.0, 4.6], 10.2)

# ---- 1.3b payment states ----
s = S()
head(s, "1.3b · payment states", "Failed vs Stuck vs Abandoned",
     "The definition L2 corrected once already — and the one L1 will get wrong if we do not nail it.")
root = node(s, 0.9, 2.6, 2.4, 0.85, "Payment not success", "what happened?", WHITE, INK)
f1 = node(s, 4.4, 2.05, 2.6, 0.8, "FAILED", "gateway declined", REDBG, RED, tcolor=RED)
f2 = node(s, 4.4, 3.05, 2.6, 0.8, "STUCK", "commit response present", AMBERBG, AMBER, tcolor=AMBER)
f3 = node(s, 4.4, 4.05, 2.6, 0.8, "ABANDONED", "no commit response", PANEL, FAINT, tcolor=MUTED)
for t in (f1, f2, f3): arrow_r(s, root, t)
for t, txt in ((f1, "Customer was NOT charged. Read the decline reason, advise, retry."),
               (f2, "Customer MAY have been charged. Escalate — do not tell them to retry."),
               (f3, "Customer never finished. Normal behaviour. Not an incident.")):
    text(s, 7.3, t[1] + 0.22, 5.4, 0.5, txt, 11, BODY, spacing=1.2)
code(s, 0.62, 5.15, 12.1, 0.95,
     "-- the single field that decides it (mirrors Payment#actual_pending? in selfcare-backend)\n"
     "lower(status) IN ('pending','initiated')\n"
     "  AND payment_commit_response IS NOT NULL AND payment_commit_response::text NOT IN ('','{}','null')", 10)
callout(s, 0.62, 6.3, 12.1, 1.0, "why this cost us",
        "Before the fix, 473 abandoned checkouts were reported as stuck. Get it wrong in the other direction "
        "and you strand a customer who really was charged. This is the definition to argue about today.", "warn", 11.5)

# ---- 1.4 attribution ----
s = S()
head(s, "1.4 · attribution", "The trap that made an unlaunched channel look live",
     "Commercial said soob had not launched. The console said 311 orders. The console was wrong.")
lane(s, 0.62, 2.42, 5.95, 2.4, "WRONG — plan_channels (config)", REDBG, RED, RED, label_above=True)
w0 = node(s, 1.15, 2.65, 1.5, 0.6, "1 order", None, WHITE, RED)
for i in range(4):
    n_ = node(s, 3.5, 2.55 + i * 0.5, 2.6, 0.4, ["tygo", "soob", "posa", "apollo"][i], None, REDBG, RED, tsize=9.5)
    arrow(s, 2.7, 2.95, 3.45, 2.75 + i * 0.5, RED, 1.1)
text(s, 0.85, 4.42, 5.5, 0.24, "one order counted under EVERY enabled channel", 10, RED, bold=True)

lane(s, 6.77, 2.42, 5.95, 2.4, "RIGHT — external_service_name (fact)", GREENBG, GREEN, GREEN, label_above=True)
r0 = node(s, 7.3, 2.95, 1.5, 0.6, "1 order", None, WHITE, GREEN)
r1 = node(s, 9.9, 2.95, 2.3, 0.6, "tygo", None, GREENBG, GREEN)
arrow_r(s, r0, r1, GREEN)
text(s, 7.0, 4.42, 5.5, 0.24, "stamped on the row at creation — exactly one", 10, GREEN, bold=True)

table(s, 0.62, 5.05, 12.1,
      [["Channel", "Reported (wrong)", "Actual", "Why it mattered"],
       ["soob", "311 orders", "65 — all test", "we nearly told commercial an unlaunched channel was selling"],
       ["tygo", "—", "26,107", "the real reseller volume was buried in the noise"]],
      [2.2, 3.0, 2.4, 4.5], 10.5)
callout(s, 0.62, 6.3, 12.1, 0.95, "the rule that came out of it",
        "Attribute from the fact stamped on the row, never from a table describing what is POSSIBLE. "
        "Four more defects of this shape were found in the same audit — if a number looks too good, say so.", "key", 11.5)

# ---- 1.5 coverage ----
s = S()
head(s, "1.5 · coverage", "What the console can and cannot see",
     "A blind spot you know about is a gap. One you do not know about is a false sense of safety.")
for i, (t, items, c, bgc) in enumerate([
    ("WE SEE", ["15 tables from the replica", "orders · payments · activation", "eligibility · nafath · delivery",
                "API GW reachability (probe)", "OSB read-path log (new this week)"], GREEN, GREENBG),
    ("WE DO NOT", ["BSS / Oracle BRM app logs", "OSB / WebLogic server logs", "Semati / TCC provider logs",
                   "firewall & SOC logs", "UPG / HyperPay gateway logs"], RED, REDBG)]):
    x = 0.62 + i * 6.15
    box(s, x, 2.15, 5.95, 2.75, bgc, c)
    text(s, x + 0.3, 2.35, 5.3, 0.26, t, 11, c, bold=True)
    for j, it in enumerate(items):
        text(s, x + 0.3, 2.75 + j * 0.4, 5.4, 0.36, "· " + it, 11, BODY)
callout(s, 0.62, 5.1, 12.1, 1.0, "the honest sentence",
        "The console sees the digital channel's OWN database, plus a few optional feeds. It is not a log "
        "aggregator. Treating it as one is exactly the assumption that let INC0016809 run for an hour.", "warn", 11.5)
verdict(s, 0.62, 6.3, 12.1, 0.95, "review", "Which missing source would change your day the most?",
        ["BSS app logs", "Semati provider", "Gateway logs", "None — this is enough"])

# ---- REVIEW BLOCK 1 ----
s = S()
head(s, "session 1 · review block", "Sign off the definitions",
     "Nothing here is settled until you say it is. A veto today is far cheaper than a wrong number in front of a customer.")
y = checklist(s, 0.62, 2.2, 6.0, [
    "Eligibility pass matches how CITC decides",
    "Total payment counts the right population",
    "Pending → stuck threshold agreed (currently 3h)",
    "Physical vs eSIM split is complete",
    "Not-assigned excludes partners correctly",
    "Partner list is complete (tygo, soob, …?)",
    "Delivered state list is current"], "definitions confirmed")
checklist(s, 6.9, 2.2, 5.8, [
    "Activated agrees with BSS",
    "10 / 45 min freshness thresholds accepted",
    "Attribution source agreed (never plan_channels)",
    "Coverage gaps understood and accepted",
    "Someone owns each disputed definition"], "and")
box(s, 0.62, 5.5, 12.1, 0.95, PANEL, LINE)
text(s, 0.92, 5.68, 11.6, 0.24, "HOMEWORK BEFORE SESSION 2", 9.5, GREEN, bold=True)
text(s, 0.92, 5.98, 11.6, 0.42, "Take FIVE real cases you handled today. Find each one in the console. Log every place the "
     "console disagreed with what you knew, or could not answer the question at all.", 11.5, BODY, spacing=1.2)
verdict(s, 0.62, 6.6, 12.1, 0.85, "before we move on", "Do you trust the numbers enough to alert on them?",
        ["Yes — go to alerts", "Not yet — fix first"])

# ══════════════════════════════════ SESSION 2 DIVIDER ══════════════════════════════════
s = S(DARK)
section(s, "session 2 · day 2", "Are the alerts right?",
        f"All {N_RULES} rules, every threshold, and the incident that got through anyway",
        [("2.1", "How a rule actually evaluates — window, sample, dimension", ""),
         ("2.2", f"The {N_RULES} rules by severity — what pages you at 3am", ""),
         ("2.3", "Baselines vs absolutes — the 55% problem", ""),
         ("2.4", "INC0016809 — how a P1 ran for an hour with a green console", ""),
         ("2.5", "Correlation, escalation, SLA & error budget", ""),
         ("", "REVIEW BLOCK — sign off every threshold", "")])

# ---- 2.1 how a rule evaluates ----
s = S()
head(s, "2.1 · mechanics", "How a rule actually evaluates",
     "Four things must all be true. Most 'why didn't it fire?' answers are one of these four.")
steps = [("METRIC", "a value + a sample size,\ncomputed per window", BLUE),
         ("WINDOW", "0.25h … 24h\nthe rule's own slice", PURPLE),
         ("MIN SAMPLE", "enough data to\nmean anything", AMBER),
         ("OPERATOR", "gte / lte\nvs threshold", GREEN)]
for i, (t, b, c) in enumerate(steps):
    x = 0.62 + i * 3.1
    box(s, x, 2.3, 2.75, 1.35, WHITE, c)
    text(s, x, 2.5, 2.75, 0.26, t, 11.5, c, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, 2.85, 2.75, 0.6, b, 10, BODY, align=PP_ALIGN.CENTER, spacing=1.2)
    if i < 3: arrow(s, x + 2.8, 2.97, x + 3.05, 2.97, FAINT)
code(s, 0.62, 3.95, 12.1, 0.75,
     "fired = value != null  AND  sample >= min_sample  AND  OP(value, threshold)  AND  within active hours", 11)
table(s, 0.62, 5.05, 12.1,
      [["Symptom", "Almost always", "Where to look"],
       ["Rule never fires", "min_sample never reached at quiet hours", "the rule's sample vs your real volume"],
       ["Rule fires nightly", "threshold set below the natural baseline", "the metric chart, not the rule"],
       ["Fired but nobody acted", "no runbook, or wrong team", "the runbook on the alert itself"],
       ["Ten alerts, one cause", "correlation not configured for that pair", "session 2.5"]],
      [3.2, 4.4, 4.5], 10.5)
banner(s, 0.62, 6.85, 12.1, "min_sample is the most common reason a real incident goes unnoticed. Ask it of every rule today.", "warn", 11.5)

# ---- 2.2 rules tour ----
s = S()
head(s, "2.2 · the rule book", f"{N_RULES} rules — what actually pages you",
     "Severity is a promise about response. If a P1 does not deserve a phone call, it is not a P1.")
for i, (sev, n, meaning, c, bgc) in enumerate([
    ("P1", N_P1, "Page now. Customers are blocked or money is at risk.", RED, REDBG),
    ("P2", N_P2, "Same shift. Degraded, not stopped.", AMBER, AMBERBG),
    ("P3", N_P3, "Watch. Early warning, review in hours.", BLUE, BLUEBG)]):
    x = 0.62 + i * 4.1
    box(s, x, 2.15, 3.9, 1.15, bgc, c)
    text(s, x + 0.3, 2.32, 1.0, 0.4, sev, 20, c, bold=True)
    text(s, x + 1.25, 2.35, 0.9, 0.34, str(n), 20, INK, bold=True)
    text(s, x + 0.3, 2.82, 3.4, 0.4, meaning, 10, BODY, spacing=1.15)
table(s, 0.62, 3.55, 12.1,
      [["P1 rule", "Fires when", "Your call today"],
       ["bss_soap_fault_1500", "≥5 SOAP faults / 30 min", "is 5 right, or too twitchy?"],
       ["apigw_node_down", "a known-good GW node stops answering", "should 1 node be P1, or only 2+?"],
       ["upg_hard_down", "zero UPG successes in 30 min", "does failover make this survivable?"],
       ["semati_provider_down", "≥50% provider errors / 15 min", "matches the TCC pattern you see?"],
       ["payment_duplicate", "≥25 suspected duplicates / 6h", "is 25 the right pain threshold?"],
       ["nafath_fail_storm", "≥50% Nafath failures / 1h", "50% or lower?"],
       ["citc_upstream_down", "Semati AND Nafath both degraded", "does one P1 beat two?"]],
      [3.4, 4.2, 4.5], 10.2)

# ---- 2.3 baselines ----
s = S()
head(s, "2.3 · baselines", "The 55% problem",
     "Semati fails about 55% on a completely normal day. Any fixed threshold is wrong.")
box(s, 0.62, 2.2, 5.95, 2.4, REDBG, RED)
text(s, 0.92, 2.4, 5.3, 0.26, "FIXED THRESHOLD", 10.5, RED, bold=True)
for j, t in enumerate(["set at 10% → fires every day, forever",
                       "set at 90% → sleeps through a real outage",
                       "either way it gets muted within a week"]):
    text(s, 0.92, 2.78 + j * 0.42, 5.4, 0.38, "· " + t, 11, BODY)
text(s, 0.92, 4.1, 5.4, 0.34, "A muted rule is worse than no rule — it looks like coverage.", 11, RED, bold=True)

box(s, 6.77, 2.2, 5.95, 2.4, GREENBG, GREEN)
text(s, 7.07, 2.4, 5.3, 0.26, "BASELINE-RELATIVE", 10.5, GREEN, bold=True)
for j, t in enumerate(["compare to the same hour, recent weeks",
                       "alert on CHANGE from normal",
                       "gate with a minimum sample"]):
    text(s, 7.07, 2.78 + j * 0.42, 5.4, 0.38, "· " + t, 11, BODY)
text(s, 7.07, 4.1, 5.4, 0.34, "55% belongs on the SLA page. Alerts are for what just got worse.", 11, GREEN, bold=True)

table(s, 0.62, 4.85, 12.1,
      [["Where we used a baseline", "Why", "Confirm with us"],
       ["Semati / Nafath rates", "chronic upstream failure is normal", "is 55% still today's normal?"],
       ["API GW nodes", "2 of 4 nodes were never opened to us", "should we chase the other 2?"],
       ["Anomaly detection", "seasonal — 3am ≠ 3pm", "does the seasonal band look right to you?"],
       ["activation_fail_storm", "MIXES BSS + Semati — kept insensitive on purpose", "agree we split rather than lower it?"]],
      [3.9, 4.2, 4.0], 10.5)

# ---- 2.4 INC0016809 ----
s = S()
head(s, "2.4 · case study", "INC0016809 — a P1 ran for an hour with a green console",
     "06 Aug. The most useful thing that happened this month, because it showed exactly what we were blind to.")
timeline(s, 1.2, 3.0, 11.0, [
    (0.0, "05:00", "firewall CRQ…190352 done"),
    (0.18, "05:00–05:30", "SOAP faults begin"),
    (0.5, "06:23", "P1 raised"),
    (0.72, "07:07", "42 faults / 30 min"),
    (1.0, "07:49", "console: 3 alerts, none this")], RED)
box(s, 0.62, 3.9, 12.1, 1.0, REDBG, RED)
text(s, 0.92, 4.06, 11.6, 0.24, "THE FAULT — and it was in OUR database the whole time", 9.5, RED, bold=True)
text(s, 0.92, 4.36, 11.6, 0.45, '{"responseCode":"1500","responseMessage":"unexpected XML tag … expected '
     'createSubscriptionTransactionResponse but found: …Fault"}', 10, INK, font=MONO, spacing=1.2)
for i, (t, b) in enumerate([
    ("No rule read the response code", "A code that is NEVER normal appeared 42 times in 30 minutes and nothing evaluated it."),
    ("The one rule that could have seen it was gated out", "activation_fail_storm needs 70% on min_sample 20. Reality: 46% on n=13."),
    ("And 70% was not a mistake", "That metric mixes BSS with Semati. The threshold HAD to clear Semati's noise — which made it structurally blind.")]):
    y = 5.05 + i * 0.62
    text(s, 0.62, y, 0.35, 0.3, str(i + 1) + ".", 12, RED, bold=True)
    text(s, 1.02, y, 5.6, 0.3, t, 11.5, INK, bold=True)
    text(s, 6.9, y, 5.8, 0.5, b, 10.5, BODY, spacing=1.15)
banner(s, 0.62, 6.93, 12.1, "Four new rules came out of this. Your job today: decide whether their thresholds are right.", "key", 11.5, 0.45)

# ---- 2.4b the four new rules ----
s = S()
head(s, "2.4b · the fix", "Four new rules — thresholds set by me, owned by you",
     "I picked these numbers from one incident. You have seen hundreds. Change them.")
table(s, 0.62, 2.2, 12.1,
      [["Rule", "Sev", "Fires when", "Why this number", "Your call"],
       ["bss_soap_fault_1500", "P1", "≥5 SOAP faults / 30 min", "the fault is never normal", "5 · 10 · 20?"],
       ["bss_degraded", "P2", "≥25% BSS failure / 1h", "clean baseline once Semati excluded", "25% right?"],
       ["bss_fail_burst", "P2", "≥6 BSS failures / 30 min", "counts, so low volume still alerts", "6 · 10?"],
       ["bss_error_dominant", "P3", "one code ≥60% of failures", "one code = upstream fault", "60% right?"]],
      [3.1, 0.7, 3.0, 3.4, 1.9], 10.2)
box(s, 0.62, 4.5, 12.1, 1.35, GREENBG, GREEN)
text(s, 0.92, 4.68, 11.6, 0.26, "THE DESIGN DECISION WORTH ARGUING WITH", 9.5, GREEN, bold=True)
text(s, 0.92, 4.98, 11.6, 0.75, "bss_fail_burst COUNTS failures instead of rating them. A rate needs a denominator, and at "
     "07:00 KSA the denominator was 13 — every rate rule sat the incident out. A count has nothing to be starved of: "
     "six failures is six failures at 03:00 and at 13:00 alike.", 11.5, BODY, spacing=1.25)
verdict(s, 0.62, 6.05, 12.1, 1.15, "decision", "Are these four thresholds right for how the platform actually behaves?",
        ["Accept as-is", "Adjust today", "Watch a week, then decide"])

# ---- 2.5 correlation & escalation ----
s = S()
head(s, "2.5 · correlation, escalation, SLA", "One outage should page once",
     "And when it pages, the ladder should climb without anyone remembering to climb it.")
rootn = node(s, 0.9, 2.5, 2.6, 0.9, "Semati down", "root cause", REDBG, RED, tcolor=RED)
for i, lbl in enumerate(["eligibility fail", "activation fail", "MNP blocked", "change plan", "SIM swap", "checkout drop"]):
    yy = 2.15 + i * 0.62
    n_ = node(s, 5.4, yy, 2.9, 0.5, lbl, None, PANEL, FAINT, tsize=9.5)
    arrow(s, 3.55, 2.95, 5.35, yy + 0.25, FAINT, 1.0)
text(s, 8.6, 2.4, 4.1, 0.24, "SUPPRESSED", 10, GREEN, bold=True)
text(s, 8.6, 2.7, 4.1, 1.1, "Six symptoms, one page. Without this you get ten alerts, ten acknowledgements, "
     "and a bridge call arguing about which is real.", 11, BODY, spacing=1.25)
box(s, 8.6, 4.0, 4.12, 1.75, PANEL, LINE)
text(s, 8.85, 4.16, 3.6, 0.24, "ESCALATION LADDER", 9.5, GREEN, bold=True)
for j, t in enumerate(["fires → owner", "+15 min unacked → L2 lead", "+30 min → duty manager", "ack STOPS the climb"]):
    text(s, 8.85, 4.48 + j * 0.32, 3.7, 0.28, "· " + t, 10.5, BODY)
table(s, 0.62, 6.0, 7.7,
      [["SLA view", "Means"],
       ["Met", "within target, budget intact"],
       ["At risk", "still met — error budget nearly spent"],
       ["Breached", "target missed for the period"]],
      [2.4, 5.3], 10.5)

# ---- REVIEW BLOCK 2 ----
s = S()
head(s, "session 2 · review block", "Sign off every threshold",
     "We go through the rule list together. Each one gets a verdict — keep, change, or delete.")
table(s, 0.62, 2.15, 12.1,
      [["Verdict", "Means", "What happens next"],
       ["KEEP", "threshold matches reality", "nothing — it is now yours"],
       ["CHANGE", "right idea, wrong number", "changed live in the session, with the metric chart open"],
       ["DELETE", "will never be actionable", "disabled — better than a muted rule pretending to be coverage"],
       ["MISSING", "an incident this would not catch", "new rule drafted before session 3"]],
      [1.9, 4.3, 5.9], 10.5)
y = checklist(s, 0.62, 4.1, 6.0, [
    "All 15 P1 rules reviewed and owned",
    "The 4 new BSS thresholds accepted or changed",
    "Every noisy rule identified by name",
    "Every missing detection written down"], "by the end of the session")
checklist(s, 6.9, 4.1, 5.8, [
    "Escalation contacts confirmed per severity",
    "Correlation pairs agreed",
    "Someone owns rule tuning going forward"], "and")
box(s, 0.62, 6.15, 12.1, 0.9, PANEL, LINE)
text(s, 0.92, 6.32, 11.6, 0.24, "HOMEWORK BEFORE SESSION 3", 9.5, GREEN, bold=True)
text(s, 0.92, 6.62, 11.6, 0.38, "For every alert that fires before we next meet, mark it actionable or noise. "
     "That single column is how we tune for the next six months.", 11.5, BODY, spacing=1.2)

# ══════════════════════════════════ SESSION 3 DIVIDER ══════════════════════════════════
s = S(DARK)
section(s, "session 3 · day 3", "Is it useful in a real incident?",
        "Not a demo — three real incidents, replayed, with you driving",
        [("3.1", "Troubleshoot board — from symptom to cause", ""),
         ("3.2", "Timeline & Subscriber 360 — one customer, end to end", ""),
         ("3.3", "Replay: INC0016809 (BSS 1500)", ""),
         ("3.4", "Replay: UPG down, masked by failover", ""),
         ("3.5", "Replay: Semati flapping — alert or ignore?", ""),
         ("", "VERDICT + ranked backlog + who owns what", "")])

# ---- 3.1 troubleshoot ----
s = S()
head(s, "3.1 · troubleshoot", "From symptom to cause",
     "The board answers one question: what is failing right now, and for how many people?")
for i, (t, b) in enumerate([
    ("1 · Category tiles", "Ten failure families. The tile that is red tells you which pipeline stage broke."),
    ("2 · Live feed", "The actual failing rows — order ref, customer, gateway, error code, timestamp."),
    ("3 · Drill-down", "The vendor's own response text. What UPG or BSS actually said, not our paraphrase."),
    ("4 · Timeline", "One order, every step, with the gap where it stopped.")]):
    y = 2.2 + i * 1.15
    box(s, 0.62, y, 12.1, 1.0)
    rule(s, 0.9, y + 0.18, 0.4, GREEN, 0.03)
    text(s, 0.9, y + 0.38, 11.4, 0.26, t, 12.5, INK, bold=True)
    text(s, 0.9, y + 0.68, 11.4, 0.26, b, 10.5, BODY)
callout(s, 0.62, 6.9, 12.1, 0.45, "the test", "Can you answer 'is this one customer or an incident?' in under 60 seconds?", "key", 11.5)

# ---- 3.3 replay 1 ----
s = S()
head(s, "3.3 · replay", "INC0016809 — you drive",
     "Simulate replay against 05:00–06:00 KSA. Same data, same rules, no hindsight.")
table(s, 0.62, 2.2, 12.1,
      [["Step", "What you do", "What we are testing"],
       ["1", "Set the window to 05:00–06:00, hit Simulate replay", "does the P1 fire, and when?"],
       ["2", "Open the alert — read the runbook on it", "is the runbook enough to act without asking anyone?"],
       ["3", "Troubleshoot → Activation (BSS), find the dominant code", "can you get to '1500' without help?"],
       ["4", "Say the blast radius out loud", "activation + plan renew + login — is that visible?"],
       ["5", "Draft the message you'd send to the bridge", "does the console give you the numbers for it?"]],
      [1.0, 6.2, 4.9], 10.5)
verdict(s, 0.62, 4.9, 12.1, 1.0, "verdict", "Would this have saved you time on the real bridge call?",
        ["Yes, an hour", "Some", "No — here is what was missing"])
callout(s, 0.62, 6.15, 12.1, 1.05, "be harsh here",
        "A polite 'yes it's useful' is worth nothing to me. The useful answer is the specific click that "
        "was missing, the number you had to go elsewhere for, the screen you would not have opened.", "warn", 11.5)

# ---- 3.6 ownership ----
s = S()
head(s, "3.6 · after this week", "Who owns what",
     "The console stops being mine at the end of session 3. That is the point of doing L2 first.")
table(s, 0.62, 2.2, 12.1,
      [["Area", "Owner after this week", "Cadence"],
       ["Daily health — is the data fresh and sane?", "L2 on duty", "each morning, 2 minutes"],
       ["Alert tuning — thresholds, noise, gaps", "L2 lead + console owner", "weekly, with the metric charts open"],
       ["New detections after an incident", "whoever ran the incident", "written up within 48h"],
       ["Definitions — what a KPI counts", "L2 agree, owner implements", "on change"],
       ["Training L1", "L2, once this review closes", "next series"]],
      [4.6, 4.0, 3.5], 10.5)
y = checklist(s, 0.62, 4.5, 12.1, [
    "Three things we agreed to change this week, with names against them",
    "The ranked backlog of what is still missing",
    "Who opens the console each morning, and what they check",
    "When we review the rules again, in the calendar before we leave",
    "Green light (or not) for the L1 workshops"], "before anyone leaves the room")
verdict(s, 0.62, 6.75, 12.1, 0.65, "the only question that matters", "Is this ready for L1?", ["Yes", "Yes, after these fixes", "No"])

# ══════════════════════════════════ WHAT I NEED ══════════════════════════════════
s = S()
head(s, "the ask", "What I need from you this week",
     "Short version: use it for real, and tell me every single thing that is wrong.")
for i, (t, b, c, bgc) in enumerate([
    ("Be specific, not polite", "'The payment page is confusing' cannot be fixed. 'Decline reason is not visible "
     "without two clicks' was fixed the same day.", RED, REDBG),
    ("Report numbers that look wrong FIRST", "Every serious defect so far was found by someone saying 'that looks "
     "wrong'. The reseller figures. The staleness banner. The stuck payments. All of them.", GREEN, GREENBG),
    ("Say when an alert wasted your time", "Noise is not a minor complaint — it is how a monitoring system dies. "
     "A rule you ignore is worse than one that does not exist.", AMBER, AMBERBG)]):
    y = 2.15 + i * 1.5
    box(s, 0.62, y, 12.1, 1.35, bgc, c)
    text(s, 0.92, y + 0.2, 11.5, 0.28, t, 13.5, INK, bold=True)
    text(s, 0.92, y + 0.56, 11.5, 0.65, b, 11.5, BODY, spacing=1.25)
box(s, 0.62, 6.75, 12.1, 0.6, GREENBG, GREEN)
text(s, 0.95, 6.9, 11.5, 0.32, "By Friday I want a ranked list of what is wrong — not a list of what is good.", 13, GREEN, bold=True)

# ══════════════════════════════════ CLOSE ══════════════════════════════════
s = S(DARK)
text(s, 0, 2.6, 13.333, 0.3, "SESSION 1 STARTS TODAY", 11, GREEN, bold=True, align=PP_ALIGN.CENTER)
text(s, 0, 3.0, 13.333, 0.75, "Bring a real case", 34, WHITE, bold=True, align=PP_ALIGN.CENTER)
rule(s, 6.15, 3.95, 1.05, GREEN)
text(s, 0, 4.25, 13.333, 0.55, "One you handled this week, that was harder than it should have been.\nWe will find it in the console together.",
     14, RGBColor(0xE2,0xE8,0xF0), align=PP_ALIGN.CENTER, spacing=1.35)
text(s, 0, 5.4, 13.333, 0.3, CFG.get("console_url", ""), 12, RGBColor(0x94,0xA3,0xB8), align=PP_ALIGN.CENTER, font=MONO)

prs.save(OUT)
print(f"✓ {OUT}")
print(f"  slides: {len(prs.slides._sldIdLst)}")
