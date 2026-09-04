#!/usr/bin/env python3
"""
Session-2 (FINAL) quiz — Kahoot-compatible import (also works as the source for Slido /
Mentimeter / Microsoft Forms). Questions MATCH the deck slides 1:1 (4 warm-up + 8 final),
shortened to Kahoot's limits (question ≤120 chars, answers ≤75, times from the allowed set).

USAGE   python3 build_l2_quiz_import.py
OUTPUT  Salam-Console-L2-Session2-Quiz.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Salam-Console-L2-Session2-Quiz.xlsx")

# (question, [a1..a4], correct_1based, seconds, round_label)
Q = [
# ── WARM-UP · session-1 recap ──
("Dashboard says “updated 4 min ago”; prod had an order 2 min ago. Broken?",
 ["Yes — must be real-time", "No — 4–7 min lag is healthy steady state", "Only if lag >15 min", "It's cached from midnight"], 2, 30, "Warm-up"),
("Subscriber 360 shows 05*****290 to an L1 agent. Why?",
 ["Display bug", "Customer opted out", "PII masked per role, server-side", "Number ported out"], 3, 30, "Warm-up"),
("Orders KPI and the order tree show “✓ reconciles”. That proves…",
 ["Tree is drawn from the KPI", "Two independent counts MATCH", "Nothing — decorative", "Rollups are enabled"], 2, 30, "Warm-up"),
("Rule: min_sample 20. Window has 3 events, 100% failing. Result?",
 ["Fires — 100% is critical", "Fires as P3", "Nothing — min_sample gates noise", "Pages on-call directly"], 3, 30, "Warm-up"),
# ── FINAL · both sessions ──
("Paid physical reseller order, courier required, NO delivery row after 30 min. Which box?",
 ["Not Assigned", "Shop pickup (reseller)", "Courier not created", "Partner-fulfilled"], 3, 30, "Final"),
("apollo_require_delivery = false on a tygo order means…",
 ["Courier delivery required", "Delivery skipped — collect in shop", "Order not eligible", "eSIM only"], 2, 30, "Final"),
("Pending payment, NO gateway commit response, 2h old. It is…",
 ["Stuck — page L2", "Duplicate suspect", "Abandoned checkout — normal", "ZATCA failure"], 3, 30, "Final"),
("OSB READ-path faults (get-subscription-profile) live in…",
 ["activation_logs (replica)", "dms_audit_logs.uil_logs (MySQL)", "payments.extra", "Not recorded"], 2, 30, "Final"),
("Semati returns 715 “Service not available” on 60% of calls. Failing layer?",
 ["Transport / TLS", "Provider APP — up but refusing", "Our replica sync", "API gateway"], 2, 30, "Final"),
("Yusr header: “Data-only mode · LLM offline”. You get…",
 ["Nothing — bot is down", "Cached answers", "Correct raw data, no AI reasoning", "Cloud fallback answers"], 3, 30, "Final"),
("Yusr's case memory (saved on 👍)…",
 ["Stores full MSISDNs", "Scrubs identifiers before storage", "Saves super_admin only", "Uploads to Ollama.com"], 2, 30, "Final"),
("A dashboard number looks impossible. First move?",
 ["Screenshot to group chat", "Assume broken, use Grafana", "Drill in, verify with sql.cjs", "Restart the console"], 3, 30, "Final"),
]

# validate Kahoot limits
ALLOWED = {5,10,20,30,60,90,120,240}
for q,a,c,t,r in Q:
    assert len(q) <= 120, f"question too long ({len(q)}): {q}"
    for x in a: assert len(x) <= 75, f"answer too long ({len(x)}): {x}"
    assert 1 <= c <= 4 and t in ALLOWED

wb = Workbook(); ws = wb.active; ws.title = "Kahoot import"
hdr = ["Question - max 120 characters", "Answer 1 - max 75 characters", "Answer 2 - max 75 characters",
       "Answer 3 - max 75 characters", "Answer 4 - max 75 characters",
       "Time limit (sec) - 5,10,20,30,60,90,120 or 240", "Correct answer(s) - choose at least one"]
ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E7F8EF")
for q,a,c,t,r in Q: ws.append([q]+a+[t,c])
widths = [70,32,32,32,32,14,12]
for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width = w
# answer-key sheet for the facilitator
ks = wb.create_sheet("Answer key")
ks.append(["#","Round","Question","Correct"]);
for c in ks[1]: c.font = Font(bold=True)
for i,(q,a,c,t,r) in enumerate(Q,1): ks.append([i,r,q,"ABCD"[c-1]+" — "+a[c-1]])
ks.column_dimensions["C"].width = 70; ks.column_dimensions["D"].width = 50
wb.save(OUT)
print("saved:", OUT, f"· {len(Q)} questions (4 warm-up + 8 final)")
